"""
revise_report_v4.py

Takes the current report (BCSE497J_Project_I_Report_Final.docx) and produces
a corrected version, applying every fix as an explicit, re-runnable step so
the change set is auditable rather than hand-made.

What this script fixes, grouped by the reason it is a fix:

  A. Document defects
     - two leftover TEMPLATE SAMPLE images (a generic Gantt chart, and an
       architecture diagram of a Twitter/GPS mobile app) sitting directly
       above this project's own figure captions;
     - a duplicate section number: the body had "3.2 Feasibility Study" AND
       "3.2 System Specification" (with 3.2.1/3.2.2 appearing twice), while
       the Table of Contents already said 3.3;
     - Table-of-Contents labels that disagreed with the body;
     - a reference list split across four separately-numbered Word lists, so
       the rendered numbering restarted three times and no in-text citation
       [n] pointed at the right entry;
     - references [2]-[6] listed but never cited anywhere in the body;
     - figure captions missing for Figs. 3-6;
     - a requirements sentence stranded in the middle of the literature
       review;
     - leftover template instruction text and "/** Sample **/" markers.

  B. Claims that the code does not support
     - "Performance: ... in real-time with minimal latency" (template sample
       text: the system runs a batch simulation and plays back a recorded
       trajectory; nothing is real-time);
     - "Security: ... strong access controls and encryption mechanisms"
       (template sample text: there is no authentication, no accounts and no
       encryption anywhere in the project, and section 3.3.2 says so);
     - "unconstrained LLM-generated parallel code is documented to be
       unreliable" -- asserted as documented, but none of the cited works
       support it;
     - an understated storage requirement that ignored the size of the local
       language model the framework proposes.

  C. Diagrams
     - all six figures replaced with the v4 set (see make_report_diagrams_v4.py).

Run:  python deliverables/revise_report_v4.py
"""

import os
import re
import shutil
import sys

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
FIGURES = os.path.join(HERE, "figures")

SRC = r"C:\Users\SHIVANI\Downloads\BCSE497J_Project_I_Report_Final.docx"
OUT = r"C:\Users\SHIVANI\Downloads\BCSE497J_Project_I_Report_Final_v2.docx"


# ---------------------------------------------------------------------------
# Small helpers for editing an existing document in place
# ---------------------------------------------------------------------------
def paras(doc):
    return doc.paragraphs


def find(doc, prefix, start=0):
    """Index of the first paragraph whose text starts with `prefix`."""
    for i, p in enumerate(doc.paragraphs[start:], start):
        if p.text.strip().startswith(prefix):
            return i
    raise LookupError(f"paragraph not found: {prefix!r}")


def find_exact(doc, text, start=0):
    for i, p in enumerate(doc.paragraphs[start:], start):
        if p.text.strip() == text:
            return i
    raise LookupError(f"paragraph not found (exact): {text!r}")


def delete_paragraph(p):
    p._p.getparent().remove(p._p)


def set_text(p, text):
    """Replaces a paragraph's content with `text`, keeping the formatting of
    its first run (so headings stay headings, list items stay list items)."""
    if not p.runs:
        p.add_run(text)
        return
    p.runs[0].text = text
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)


def set_labelled(p, label, body):
    """Rewrites a 'Label: body text' bullet, keeping the label bold like the
    surrounding bullets in the same list."""
    if not p.runs:
        p.add_run("")
    first = p.runs[0]
    first.text = label
    first.bold = True
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)
    rest = p.add_run(body)
    rest.bold = False
    rest.font.name = first.font.name
    rest.font.size = first.font.size


def copy_para_format(src_p, dst_p):
    """Clones a paragraph's pPr (indent/spacing/alignment/list membership)."""
    src_pPr = src_p._p.pPr
    if src_pPr is None:
        return
    import copy
    new_pPr = copy.deepcopy(src_pPr)
    old = dst_p._p.pPr
    if old is not None:
        dst_p._p.remove(old)
    dst_p._p.insert(0, new_pPr)


def insert_para_after(anchor_p, text="", like=None):
    """Inserts a new paragraph immediately after `anchor_p`, optionally taking
    its paragraph formatting from `like`."""
    new_p = OxmlElement("w:p")
    anchor_p._p.addnext(new_p)
    para = docx.text.paragraph.Paragraph(new_p, anchor_p._parent)
    if like is not None:
        copy_para_format(like, para)
    if text:
        run = para.add_run(text)
        if like is not None and like.runs:
            run.font.name = like.runs[0].font.name
            run.font.size = like.runs[0].font.size
    return para


def strip_list_numbering(p):
    """Removes a paragraph from its auto-numbered list (used for the
    reference list, which is numbered explicitly as [1], [2], ... instead)."""
    pPr = p._p.pPr
    if pPr is not None and pPr.numPr is not None:
        pPr.remove(pPr.numPr)


def add_bookmark(p, name, bid):
    """Wraps a paragraph in a Word bookmark so the Table of Contents can point
    a PAGEREF field at it."""
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bid))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bid))
    p._p.insert(0, start)
    p._p.append(end)


def set_pageref(cell, bookmark):
    """Replaces a Table-of-Contents page-number cell with a live PAGEREF
    field, so Word computes the real page number instead of us guessing one."""
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def r(child):
        run = OxmlElement("w:r")
        run.append(child)
        p._p.append(run)

    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    r(begin)

    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" PAGEREF {bookmark} \\h "
    r(instr)

    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    r(sep)

    t = OxmlElement("w:t")
    t.text = "1"
    r(t)

    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    r(end)


def clean_run_formatting(doc):
    """Strips the editing and web-paste artefacts the document had picked up:
    yellow/cyan/green highlighter marks left over from drafting, a grey
    paragraph background and a Roboto typeface that came in with pasted text,
    and an off-black text colour. None of these are deliberate design; they
    just make the report look like a working draft."""
    removed = {"highlight": 0, "shading": 0, "font": 0, "colour": 0}

    def scrub_run(r):
        rPr = r._r.rPr
        if rPr is None:
            return
        for tag, key in (("w:highlight", "highlight"), ("w:shd", "shading")):
            el = rPr.find(qn(tag))
            if el is not None:
                rPr.remove(el)
                removed[key] += 1
        fonts = rPr.find(qn("w:rFonts"))
        if fonts is not None:
            rPr.remove(fonts)          # inherit the document's own typeface
            removed["font"] += 1
        colour = rPr.find(qn("w:color"))
        if colour is not None:
            rPr.remove(colour)
            removed["colour"] += 1

    def scrub_para(p):
        pPr = p._p.pPr
        if pPr is not None:
            shd = pPr.find(qn("w:shd"))
            if shd is not None:
                pPr.remove(shd)
                removed["shading"] += 1
        for r in p.runs:
            scrub_run(r)

    for p in doc.paragraphs:
        scrub_para(p)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    scrub_para(p)
    return removed


def bold_lead_in(p, lead):
    """Bolds the opening phrase of a paragraph, so the literature review's
    themes are visible at a glance instead of buried in running text."""
    text = p.text
    if not text.startswith(lead):
        return False
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    a = p.add_run(lead)
    a.bold = True
    p.add_run(text[len(lead):])
    return True


def normalise_body_text(doc, first_heading="1. INTRODUCTION"):
    """Gives every ordinary body paragraph in the report the same treatment:
    justified, 1.15 line spacing, a little space after. Headings, list items,
    figures and captions are left alone."""
    started = False
    fixed = 0
    for p in doc.paragraphs:
        text = p.text.strip()
        if text.startswith(first_heading):
            started = True
        if not started or not text:
            continue
        if p._p.findall(".//" + qn("a:blip")):
            continue
        pPr = p._p.pPr
        if pPr is not None and pPr.numPr is not None:
            continue                                  # bullet / numbered item
        if text.startswith("Fig. ") or text.startswith("["):
            continue                                  # caption / reference
        runs = [r for r in p.runs if r.text.strip()]
        if runs and runs[0].font.size is not None:
            continue                                  # a sized heading
        if p.alignment == WD_ALIGN_PARAGRAPH.CENTER:
            continue
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf = p.paragraph_format
        pf.line_spacing = 1.15
        pf.space_after = Pt(8)
        fixed += 1
    return fixed


def normalise_bullets(doc, first_heading="1. INTRODUCTION"):
    """Every bulleted requirement, feasibility point and specification in this
    report is written as 'Label: explanation'. Some lists had the label bold
    and some did not, purely by accident of how they were pasted in. This
    makes them all consistent, and justifies them like the body text."""
    started = False
    fixed = 0
    label_re = re.compile(r"^([A-Z][A-Za-z0-9 ()/&+.-]{2,44}):\s")
    for p in doc.paragraphs:
        text = p.text
        if text.strip().startswith(first_heading):
            started = True
        if not started:
            continue
        pPr = p._p.pPr
        if pPr is None or pPr.numPr is None:
            continue
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.15
        m = label_re.match(text)
        if not m:
            continue
        label = m.group(1)
        first = p.runs[0] if p.runs else None
        if first is None:
            continue
        name, size = first.font.name, first.font.size
        for r in list(p.runs):
            r._r.getparent().remove(r._r)
        a = p.add_run(label)
        a.bold = True
        a.font.name, a.font.size = name, size
        b = p.add_run(text[len(label):])
        b.bold = False
        b.font.name, b.font.size = name, size
        fixed += 1
    return fixed


def keep_with_next(p):
    """Stops a heading from being stranded at the foot of a page."""
    pPr = p._p.get_or_add_pPr()
    if pPr.find(qn("w:keepNext")) is None:
        pPr.append(OxmlElement("w:keepNext"))


def page_break_before(p):
    """Starts a chapter on a fresh page, the way a bound report reads."""
    pPr = p._p.get_or_add_pPr()
    existing = pPr.find(qn("w:pageBreakBefore"))
    if existing is None:
        el = OxmlElement("w:pageBreakBefore")
        pPr.append(el)


def set_roman_front_matter(doc):
    """The front matter (cover, abstract, contents) is a separate Word
    section from the body, and the body already restarts at page 1. Numbering
    the front matter in lower-case roman is what makes the Table of Contents
    read correctly -- otherwise 'Abstract ... 2' sits above
    'INTRODUCTION ... 1'."""
    sect = doc.sections[0]
    sectPr = sect._sectPr
    pg = sectPr.find(qn("w:pgNumType"))
    if pg is None:
        pg = OxmlElement("w:pgNumType")
        sectPr.append(pg)
    pg.set(qn("w:fmt"), "lowerRoman")
    pg.set(qn("w:start"), "1")


def replace_picture(p, image_path, width_in):
    """Swaps the image inside a picture-bearing paragraph for a new file."""
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(image_path, width=Inches(width_in))


def caption(anchor_p, text, like):
    """Adds a centred figure caption below `anchor_p`."""
    cap = insert_para_after(anchor_p, like=like)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cap.add_run(text)
    run.bold = False
    run.italic = False
    return cap


# ---------------------------------------------------------------------------
def main():
    doc = docx.Document(SRC)

    # =======================================================================
    # 0. Strip drafting and web-paste artefacts (highlighter marks, a grey
    #    paragraph background, a pasted Roboto typeface, off-black text).
    # =======================================================================
    stats = clean_run_formatting(doc)
    print("[0] cleaned formatting artefacts: "
          + ", ".join(f"{v} {k}" for k, v in stats.items() if v))

    # =======================================================================
    # 1. Remove template scaffolding: formatting notes, sample markers and
    #    bracketed guidance. These are instructions to the report's author,
    #    not content, and a submitted report should not contain them.
    # =======================================================================
    scaffold_patterns = [
        r"^\(Times New Roman.*\)$",
        r"^<Contents, Times New Roman.*>$",
        r"^/\*+ *Sample *\*+/$",
        r"^< ?This section of a project report provides.*>$",
        r"^< ?Maximum of 50 journal papers are required\.>$",
        r"^<To identify the areas or aspects where existing research.*>$",
        r"^< ?To clearly define what the project aims to achieve.*>$",
        r"^< ?Clearly define the issue or challenge.*>$",
        r"^<A Gantt chart helps in organizing tasks.*>$",
        r"^<Please include appropriate hardware & software.*>$",
        r"^<Please include whichever diagram is applicable.*>$",
        r"^One page and not exceeding 300 words$",
    ]
    removed = 0
    for p in list(doc.paragraphs):
        txt = p.text.strip()
        if any(re.match(pat, txt) for pat in scaffold_patterns):
            delete_paragraph(p)
            removed += 1
    print(f"[1] removed {removed} template-instruction paragraphs")

    # =======================================================================
    # 2. ABSTRACT -- state the evidence for the correctness claim.
    # =======================================================================
    i = find(doc, "As a working demonstration of this concept")
    p = doc.paragraphs[i]
    set_text(p, p.text.rstrip().rstrip(".") +
             ". Its physical correctness is confirmed by an automated test "
             "suite that checks energy conservation directly rather than "
             "assuming it.")
    print("[2] abstract: correctness claim tied to the test suite")

    # =======================================================================
    # 3. SCOPE -- state the prototype's actual, code-enforced limits.
    # =======================================================================
    i = find(doc, "It further covers a simulation engine implementing")
    p = doc.paragraphs[i]
    set_text(p, p.text.rstrip() +
             " The prototype is scoped to systems of up to 500 particles and "
             "run lengths of up to 20 reduced time units — limits enforced by "
             "the validation layer itself — which is sufficient to exercise "
             "every stage of the pipeline on ordinary CPU hardware.")
    print("[3] scope: added the prototype's enforced limits")

    # =======================================================================
    # 4. LITERATURE REVIEW
    #    (a) remove a requirements sentence stranded inside the review;
    #    (b) add the paragraph that actually cites references [2]-[6], which
    #        were listed in the reference list but cited nowhere.
    # =======================================================================
    i = find(doc, "Accessibility: The system should not require the user")
    delete_paragraph(doc.paragraphs[i])
    print("[4a] removed the misplaced 'Accessibility' sentence from 2.1")

    i = find(doc, "Foundations of molecular dynamics.")
    anchor = doc.paragraphs[i]
    new_p = insert_para_after(anchor, like=anchor)
    new_p.add_run(
        "Production molecular dynamics software and GPU acceleration. "
        "The performance side of this problem is well solved. Plimpton's "
        "short-range parallel algorithms [2] underpin LAMMPS, and GROMACS [3] "
        "demonstrates multi-level parallelism from laptops to supercomputers. "
        "Stone et al. [4] showed early on that graphics processors suit the "
        "pairwise force calculation at the centre of molecular dynamics, and "
        "HOOMD-blue [5] and OpenMM [6] are mature packages built around GPU "
        "execution. Together these establish two things this project relies "
        "on: that GPU acceleration of the O(N²) force loop is a solved and "
        "well-precedented engineering problem rather than a research risk, "
        "and that every one of these packages is still configured through "
        "scripts, input decks or programming interfaces — which is precisely "
        "the barrier this project addresses."
    )
    i = find(doc, "Language models, structured extraction")
    set_text(doc.paragraphs[i], doc.paragraphs[i].text.replace(
        "and related structured-output/function-calling techniques",
        "and the structured-output / function-calling interfaces now standard "
        "in production language-model APIs [14]"))

    for lead in ("Foundations of molecular dynamics.",
                 "Production molecular dynamics software and GPU acceleration.",
                 "Machine learning inside molecular simulation.",
                 "Language models, structured extraction, and tool-restricted "
                 "scientific agents."):
        try:
            bold_lead_in(doc.paragraphs[find(doc, lead)], lead)
        except LookupError:
            print(f"    ! literature lead-in not found: {lead[:40]}")
    print("[4b] added the literature paragraph citing [2]-[6]; "
          "theme lead-ins emphasised")

    # =======================================================================
    # 5. RESEARCH GAP -- remove an unsupported "documented to be unreliable"
    #    claim and replace it with evidence the report can actually stand on
    #    (its own Section 1.2), plus proper citations for the named tools.
    # =======================================================================
    i = find(doc, "Mature, high-performance MD tools")
    set_text(doc.paragraphs[i],
        "Mature, high-performance MD tools — LAMMPS [2], GROMACS [3], "
        "HOOMD-blue [5] and OpenMM [6] — solve the performance problem but "
        "not the accessibility problem: all remain script- or "
        "API-configured, with no natural-language entry point. "
        "General-purpose large language models can generate simulation code "
        "directly, but nothing constrains that generated code to be "
        "physically correct, and defects in low-level parallel code tend to "
        "be silent rather than obvious — an incorrect memory-allocation "
        "cast, an unsupported double-precision atomic operation, or "
        "incomplete host/device memory management will often compile and run "
        "while producing wrong numbers, exactly as encountered first-hand in "
        "the 2024 CUDA port described in Section 1.2. Recent "
        "language-model tool-use systems in adjacent scientific domains — "
        "ChemCrow [11] and Coscientist [12] — demonstrate that restricting a "
        "language model to intent-interpretation, backed by verified tools, "
        "is a workable and trustworthy pattern; but that pattern has not "
        "been applied to GPU-accelerated molecular dynamics specifically. "
        "This project addresses that gap directly: a language-understanding "
        "component restricted to parameter extraction only, paired with a "
        "fixed, independently verifiable simulation engine, with GPU/CUDA "
        "acceleration as the framework's intended computational direction.")
    print("[5] research gap: unsupported claim replaced, citations added")

    # =======================================================================
    # 6. NON-FUNCTIONAL REQUIREMENTS -- two bullets were still the template's
    #    own sample text and contradicted the actual system.
    # =======================================================================
    i = find(doc, "Performance: The system should process and analyze data")
    set_labelled(doc.paragraphs[i], "Performance",
        ": A simulation request should return real trajectory and energy "
        "data within a few seconds for the particle counts and run lengths "
        "this prototype supports, so a user can iterate on a description "
        "interactively; the recorded trajectory should then play back in the "
        "browser at interactive frame rates.")

    i = find(doc, "Security: The system should ensure data integrity")
    set_labelled(doc.paragraphs[i], "Security and Data Handling",
        ": The system runs entirely on the user's own machine and exposes no "
        "service beyond the local loopback interface, so it holds no "
        "credentials and transmits no user data. The protection that does "
        "matter is input validation: no set of parameters, however "
        "malformed, should be able to reach the simulation engine without "
        "first being checked for physical feasibility.")
    print("[6] non-functional requirements: two template bullets corrected")

    # =======================================================================
    # 7. TECHNICAL FEASIBILITY -- add the verification evidence, so the
    #    report's correctness claims are backed by something checkable.
    # =======================================================================
    i = find(doc, "Integration: The frontend, backend, and simulation engine")
    anchor = doc.paragraphs[i]
    new_p = insert_para_after(anchor, like=anchor)
    r1 = new_p.add_run("Verification")
    r1.bold = True
    if anchor.runs:
        r1.font.name = anchor.runs[0].font.name
        r1.font.size = anchor.runs[0].font.size
    r2 = new_p.add_run(
        ": Correctness is checked rather than asserted. An automated suite of "
        "21 tests covers the validation layer, the natural-language "
        "extractor and the compiled engine, including energy-conservation "
        "checks on container-free (NVE) Lennard-Jones runs and on "
        "Hard-Sphere collisions — the standard physical test for a "
        "molecular dynamics integrator.")
    r2.bold = False
    if anchor.runs:
        r2.font.name = anchor.runs[0].font.name
        r2.font.size = anchor.runs[0].font.size
    print("[7] technical feasibility: verification bullet added")

    # =======================================================================
    # 8. Fix the duplicate 3.2 numbering in the body (the Table of Contents
    #    was already correct at 3.3 / 3.3.1 / 3.3.2).
    # =======================================================================
    for old, new in [("3.2 System Specification", "3.3 System Specification"),
                     ("3.2.1 Hardware Specification", "3.3.1 Hardware Specification"),
                     ("3.2.2 Software Specification", "3.3.2 Software Specification")]:
        i = find_exact(doc, old)
        set_text(doc.paragraphs[i], new)
    print("[8] body renumbered 3.2 -> 3.3 (System Specification)")

    # =======================================================================
    # 9. HARDWARE / SOFTWARE SPECIFICATION -- make the numbers and the named
    #    technologies match what the project actually uses and proposes.
    # =======================================================================
    i = find(doc, "Storage: Under 500 MB")
    set_labelled(doc.paragraphs[i], "Storage",
        ": Approximately 100 MB for the source code, the compiled simulation "
        "engine and the Python dependencies. A quantised local small language "
        "model, once introduced, would add roughly 0.5–4 GB depending on the "
        "model chosen.")

    i = find(doc, "Graphics Processing Unit (GPU)")
    set_labelled(doc.paragraphs[i], "Graphics Processing Unit (GPU)",
        ": An NVIDIA GPU with CUDA support, Compute Capability 6.0 (Pascal) "
        "or higher recommended — native double-precision atomicAdd() was "
        "introduced at that generation [13], and its absence on the target "
        "device was one of the concrete obstacles met in the 2024 CUDA port "
        "described in Section 1.2. This is the core computational requirement "
        "for the framework's GPU-accelerated execution, and the primary "
        "hardware its architecture is designed around.")

    i = find(doc, "Programming Languages:")
    set_labelled(doc.paragraphs[i], "Programming Languages",
        ": C++11 (simulation engine), Python 3.10 or later "
        "(backend and orchestration), and HTML/CSS/JavaScript (interface); "
        "CUDA C++ for the proposed GPU-accelerated execution backend.")

    i = find(doc, "Development Environment:")
    set_labelled(doc.paragraphs[i], "Development Environment",
        ": A standard code editor, a C++11-capable compiler (g++, via "
        "MinGW-w64 on Windows) and Python 3.10 or later. The proposed GPU "
        "backend would additionally require the NVIDIA CUDA Toolkit [13].")

    i = find(doc, "Libraries and Frameworks:")
    set_labelled(doc.paragraphs[i], "Libraries and Frameworks",
        ": FastAPI with Uvicorn for the local backend, Pydantic v2 for "
        "parameter validation, and pytest for the automated test suite. The "
        "simulation engine deliberately uses only the C++ standard library, "
        "and the browser interface uses no frontend framework — both choices "
        "keep the parts that compute or display real physics free of "
        "third-party dependencies.")
    print("[9] hardware/software specification corrected")

    # =======================================================================
    # 10. FIGURES
    # =======================================================================
    # --- Fig. 1: use the current Gantt chart, and delete the template's own
    #     sample Gantt chart, which sat directly above this project's caption.
    i = find_exact(doc, "2.5 Project Plan")
    pic_paras = [p for p in doc.paragraphs[i:i + 12]
                 if p._p.findall(".//" + qn("a:blip"))]
    assert len(pic_paras) == 2, f"expected 2 images under 2.5, found {len(pic_paras)}"
    replace_picture(pic_paras[0], os.path.join(FIGURES, "fig1_gantt_v4.png"), 6.1)
    delete_paragraph(pic_paras[1])          # template sample Gantt chart
    print("[10a] Fig. 1 replaced; template sample Gantt chart removed")

    i = find(doc, "Fig. 1.")
    set_text(doc.paragraphs[i], "Fig. 1. Project plan (Gantt chart)")
    gantt_cap = doc.paragraphs[i]
    note_p = insert_para_after(gantt_cap, like=doc.paragraphs[find(doc, "Molecular dynamics (MD) simulation is a well")])
    note_p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    note_p.add_run(
        "Figure 1 shows the plan as relative project weeks rather than "
        "calendar dates, and groups the ten phases into the five workstreams "
        "they belong to. Phases overlap where the work genuinely ran in "
        "parallel — the validation layer and the backend endpoints, for "
        "instance, were built against each other.")

    # --- Fig. 2: replace with the layered architecture diagram, and delete
    #     the template's sample architecture image (a Twitter/GPS mobile app)
    #     together with the run of empty paragraphs it left behind.
    i = find_exact(doc, "4.1 System Architecture")
    j = find(doc, "Fig. 2.", i)
    block = doc.paragraphs[i:j]
    pic_paras = [p for p in block if p._p.findall(".//" + qn("a:blip"))]
    assert len(pic_paras) == 2, f"expected 2 images under 4.1, found {len(pic_paras)}"
    replace_picture(pic_paras[0], os.path.join(FIGURES, "fig2_architecture_v4.png"), 6.2)
    delete_paragraph(pic_paras[1])          # template sample (Twitter/GPS) diagram
    for p in block:
        if p.text.strip() == "" and not p._p.findall(".//" + qn("a:blip")):
            delete_paragraph(p)
    print("[10b] Fig. 2 replaced; template sample architecture image removed")

    i = find(doc, "Fig. 2.")
    set_text(doc.paragraphs[i], "Fig. 2. System architecture")

    i = find(doc, "Figure 2 shows the proposed system architecture")
    set_text(doc.paragraphs[i],
        "Figure 2 shows the proposed system architecture as four layers "
        "inside a single machine boundary — the whole system runs locally, "
        "with no network access and no external AI service. A user's request, "
        "whether typed as a sentence or entered on the form, reaches a local "
        "language-understanding component whose role is strictly limited to "
        "extracting structured parameters; it never generates or influences "
        "simulation logic. Those parameters pass through one shared "
        "validation layer and then to a simulation-configuration stage that "
        "is backend-agnostic by design: the same configuration file can drive "
        "either a CPU execution engine or a GPU/CUDA-accelerated one without "
        "any change to the layers above it. This project stage implements and "
        "demonstrates the CPU path, shown with a solid border; GPU/CUDA "
        "execution, shown dashed, is the framework's proposed computational "
        "direction. Both paths converge on the same real trajectory and "
        "energy output, visualised identically regardless of which backend "
        "produced it.")

    # --- Figs. 3-6: replace the images, strip the template's <Mandatory> /
    #     <Optional> markers from the headings, and add the captions and
    #     explanatory paragraphs that were missing entirely.
    design_figs = [
        ("4.2.1 Data Flow Diagram", "fig3_dfd_v4.png", 6.1,
         "Fig. 3. Level-1 data flow diagram",
         "Figure 3 is a Level-1 data flow diagram in Gane–Sarson notation. "
         "The single external entity is the user; the five numbered processes "
         "are the stages a request passes through; and the three data stores "
         "are the only places data comes to rest. Two details matter. First, "
         "extraction (2.0) and validation (3.0) are separate processes: the "
         "manual input path skips extraction entirely and joins at 3.0, so "
         "both input paths are checked by exactly the same feasibility rules. "
         "Second, the diagram deliberately says nothing about how process 4.0 "
         "computes — that is the substitution point where a GPU/CUDA backend "
         "would replace the CPU one, with every flow around it unchanged."),
        ("4.2.2 Use Case Diagram", "fig4_usecase_v4.png", 5.9,
         "Fig. 4. Use case diagram",
         "Figure 4 shows the system's use cases from the single actor's point "
         "of view. Describing a simulation in natural language «include»s "
         "parameter extraction; reviewing parameters «include»s the physical "
         "feasibility check; and running a simulation «include»s execution. "
         "Execution is specialised into two alternatives — on CPU, which this "
         "project implements, and on GPU/CUDA, which it proposes — so the "
         "diagram records the intended extension without implying that it "
         "already exists."),
        ("4.2.3 Class Diagram", "fig5_class_v4.png", 6.2,
         "Fig. 5. Class diagram",
         "Figure 5 gives the class structure, with names, attributes and "
         "operations taken directly from the source files so the two can be "
         "checked against each other. The key relationship is the "
         "«interface» SimulationEngine: CPUEngine realises it today, and the "
         "proposed CUDAEngine — drawn dashed, and marked as not implemented — "
         "would realise the same interface, which is what makes the GPU "
         "backend a contained substitution rather than a redesign. Note also "
         "that the language component (ParameterExtractor) only ever creates "
         "a SimParams object; it has no relationship of any kind with the "
         "engine classes, which is the separation the whole design depends on."),
        ("4.2.4 Sequence Diagram", "fig6_sequence_v4.png", 6.2,
         "Fig. 6. Sequence diagram — running a simulation",
         "Figure 6 traces one complete run. The natural-language step is "
         "wrapped in an opt fragment because it happens only on that input "
         "path; the manual path enters at the validation call. Validation "
         "appears twice on purpose — once when the user reviews the "
         "parameters, and again inside the simulate request — so that "
         "nothing unvalidated can reach the engine even if the first check "
         "were bypassed. The engine itself is launched as a separate process "
         "against a configuration file it reads at start-up, and the loop "
         "fragment marks the per-timestep work that produces the real "
         "trajectory and energy data returned to the browser."),
    ]

    for heading_prefix, fig_file, width, cap_text, prose in design_figs:
        h = find(doc, heading_prefix)
        set_text(doc.paragraphs[h], heading_prefix)
        pic = None
        for p in doc.paragraphs[h + 1:h + 4]:
            if p._p.findall(".//" + qn("a:blip")):
                pic = p
                break
        assert pic is not None, f"no image found under {heading_prefix}"
        replace_picture(pic, os.path.join(FIGURES, fig_file), width)

        cap_p = caption(pic, cap_text, like=doc.paragraphs[h])
        cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in cap_p.runs:
            r.bold = False
            r.italic = False

        body_p = insert_para_after(cap_p)
        body_p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        body_p.paragraph_format.space_after = Pt(8)
        body_p.paragraph_format.line_spacing = 1.15
        body_p.add_run(prose)
    print("[10c] Figs. 3-6 replaced, captioned and explained")

    # =======================================================================
    # 11. A closing paragraph in 4.2 giving the actual verification results,
    #     so "physically correct" is a reported measurement, not an assertion.
    # =======================================================================
    i = find(doc, "Figure 6 traces one complete run")
    anchor = doc.paragraphs[i]
    v = insert_para_after(anchor)
    v.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    v.paragraph_format.space_after = Pt(8)
    v.paragraph_format.line_spacing = 1.15
    r1 = v.add_run("Verification of the implemented prototype. ")
    r1.bold = True
    v.add_run(
        "The design above is backed by an automated test suite of 21 tests, "
        "all passing, covering the validation layer, the natural-language "
        "extractor and the compiled engine. The two physically meaningful "
        "checks are energy-conservation tests. In a container-free "
        "(periodic, NVE) Lennard-Jones run of 100 particles, the total energy "
        "is required to drift by less than 0.5 reduced units over the whole "
        "run; in a Hard-Sphere run, where the collision arithmetic should "
        "conserve kinetic energy exactly, the tolerance is 0.05. Both hold. "
        "Energy conservation is the standard correctness criterion for a "
        "molecular dynamics integrator [1], which is why the interface "
        "presents the energy-versus-time chart alongside the 3D view: the "
        "same evidence the test suite checks is visible to the user on every "
        "run. The source code and the test suite are available in the "
        "project repository [15].")
    print("[11] verification paragraph added to 4.2")

    # =======================================================================
    # 12. REFERENCES -- rebuild as one continuous, explicitly numbered IEEE
    #     list. Previously the entries were split across four separate Word
    #     auto-numbered lists, so the rendered numbering restarted three
    #     times and no in-text citation [n] resolved to the right entry.
    # =======================================================================
    refs = [
        "D. Frenkel and B. Smit, Understanding Molecular Simulation: From "
        "Algorithms to Applications, 2nd ed. San Diego, CA, USA: Academic "
        "Press, 2002.",
        "S. Plimpton, “Fast parallel algorithms for short-range molecular "
        "dynamics,” J. Comput. Phys., vol. 117, no. 1, pp. 1–19, 1995.",
        "M. J. Abraham et al., “GROMACS: High performance molecular "
        "simulations through multi-level parallelism from laptops to "
        "supercomputers,” SoftwareX, vol. 1–2, pp. 19–25, 2015.",
        "J. E. Stone et al., “Accelerating molecular modeling applications "
        "with graphics processors,” J. Comput. Chem., vol. 28, no. 16, "
        "pp. 2618–2640, 2007.",
        "J. A. Anderson, C. D. Lorenz, and A. Travesset, “General purpose "
        "molecular dynamics simulations fully implemented on graphics "
        "processing units,” J. Comput. Phys., vol. 227, no. 10, "
        "pp. 5342–5359, 2008.",
        "P. Eastman et al., “OpenMM 7: Rapid development of high performance "
        "algorithms for molecular dynamics,” PLOS Comput. Biol., vol. 13, "
        "no. 7, e1005659, 2017.",
        "J. Behler and M. Parrinello, “Generalized neural-network "
        "representation of high-dimensional potential-energy surfaces,” "
        "Phys. Rev. Lett., vol. 98, no. 14, 146401, 2007.",
        "K. T. Butler, D. W. Davies, H. Cartwright, O. Isayev, and A. Walsh, "
        "“Machine learning for molecular and materials science,” Nature, "
        "vol. 559, no. 7715, pp. 547–555, 2018.",
        "F. Noé, A. Tkatchenko, K.-R. Müller, and C. Clementi, “Machine "
        "learning for molecular simulation,” Annu. Rev. Phys. Chem., "
        "vol. 71, pp. 361–390, 2020.",
        "J. White et al., “A prompt pattern catalog to enhance prompt "
        "engineering with ChatGPT,” arXiv:2302.11382, 2023.",
        "A. M. Bran, S. Cox, O. Schilter, C. Baldassari, A. D. White, and "
        "P. Schwaller, “Augmenting large language models with chemistry "
        "tools,” Nature Machine Intelligence, vol. 6, pp. 525–535, 2024.",
        "D. A. Boiko, R. MacKnight, B. Kline, and G. Gomes, “Autonomous "
        "chemical research with large language models,” Nature, vol. 624, "
        "no. 7992, pp. 570–578, 2023.",
        "NVIDIA Corporation, CUDA C++ Programming Guide, Release 12.5. "
        "Santa Clara, CA, USA: NVIDIA, 2024. [Online]. Available: "
        "https://docs.nvidia.com/cuda/cuda-c-programming-guide/",
        "OpenAI, “Function calling,” OpenAI Platform Documentation, 2024. "
        "[Online]. Available: "
        "https://platform.openai.com/docs/guides/function-calling",
        "S. Srivastava, “Molecular dynamics simulation of particles in a "
        "box,” GitHub repository, 2024. [Online]. Available: "
        "https://github.com/Kiwi-Supreme/"
        "Molecular-Dynamics-Simulation-of-Particles-in-a-Box",
    ]

    i = find_exact(doc, "5. REFERENCES")
    ref_anchor = doc.paragraphs[i]
    for p in list(doc.paragraphs[i + 1:]):
        delete_paragraph(p)

    style_src = doc.paragraphs[find(doc, "Molecular dynamics (MD) simulation is a well")]
    prev = ref_anchor
    for n, text in enumerate(refs, 1):
        p = insert_para_after(prev)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf = p.paragraph_format
        pf.left_indent = Inches(0.4)
        pf.first_line_indent = Inches(-0.4)
        pf.space_after = Pt(6)
        pf.line_spacing = 1.15
        run = p.add_run(f"[{n}]\t{text}")
        if style_src.runs:
            run.font.name = style_src.runs[0].font.name
            run.font.size = style_src.runs[0].font.size
        prev = p
    print(f"[12] reference list rebuilt as one continuous IEEE list "
          f"([1]–[{len(refs)}])")

    # =======================================================================
    # 13. TABLE OF CONTENTS -- correct two labels that had drifted from the
    #     body, strip the template's own formatting annotations, and replace
    #     the hand-typed page numbers with live PAGEREF fields.
    # =======================================================================
    toc = None
    for t in doc.tables:
        if t.rows and t.rows[0].cells[0].text.strip().lower().startswith("sl.no"):
            toc = t
            break
    assert toc is not None, "Table of Contents not found"

    fixes = {
        "1.2 Motivations": "1.2 Motivation",
        "3.2.2 Social Feasibility": "3.2.3 Social Feasibility",
    }
    for row in toc.rows[1:]:
        cell = row.cells[1]
        label = re.sub(r"\s*<[^>]*>", "", cell.text).strip()
        label = fixes.get(label, label)
        cell.text = ""
        para = cell.paragraphs[0]
        run = para.add_run(label)
        run.bold = label.isupper() or label == "Abstract"

    # Bookmark every heading in the body, then point the matching TOC row at it.
    heading_for_toc = {}
    body_headings = [
        ("Abstract", "ABSTRACT"),
        ("1. INTRODUCTION", "1. INTRODUCTION"),
        ("1.1 Background", "1.1 Background"),
        ("1.2 Motivation", "1.2 Motivation"),
        ("1.3 Scope of the Project", "1.3 Scope of the Project"),
        ("2. PROJECT DESCRIPTION AND GOALS", "2. PROJECT DESCRIPTION AND GOALS"),
        ("2.1 Literature Review", "2.1 Literature Review"),
        ("2.2 Research Gap", "2.2 Research Gap"),
        ("2.3 Objectives", "2.3 Objectives"),
        ("2.4 Problem Statement", "2.4 Problem Statement"),
        ("2.5 Project Plan", "2.5 Project Plan"),
        ("3. TECHNICAL SPECIFICATION", "3. TECHNICAL SPECIFICATION"),
        ("3.1 Requirements", "3.1 Requirements"),
        ("3.1.1 Functional", "3.1.1"),
        ("3.1.2 Non-Functional", "3.1.2"),
        ("3.2 Feasibility Study", "3.2 Feasibility Study"),
        ("3.2.1 Technical Feasibility", "3.2.1 Technical Feasibility"),
        ("3.2.2 Economic Feasibility", "3.2.2 Economic Feasibility"),
        ("3.2.3 Social Feasibility", "3.2.3 Social Feasibility"),
        ("3.3 System Specification", "3.3 System Specification"),
        ("3.3.1 Hardware Specification", "3.3.1 Hardware Specification"),
        ("3.3.2 Software Specification", "3.3.2 Software Specification"),
        ("4. DESIGN APPROACH AND DETAILS", "4. DESIGN APPROACH AND DETAILS"),
        ("4.1 System Architecture", "4.1 System Architecture"),
        ("4.2 Design", "4.2 Design"),
        ("4.2.1 Data Flow Diagram", "4.2.1 Data Flow Diagram"),
        ("4.2.2 Use Case Diagram", "4.2.2 Use Case Diagram"),
        ("4.2.3 Class Diagram", "4.2.3 Class Diagram"),
        ("4.2.4 Sequence Diagram", "4.2.4 Sequence Diagram"),
        ("5. REFERENCES", "5. REFERENCES"),
    ]

    bid = 900
    for toc_label, body_prefix in body_headings:
        try:
            idx = find(doc, body_prefix)
        except LookupError:
            print(f"    ! no body heading for TOC entry {toc_label!r}")
            continue
        name = f"_Toc_MD{bid}"
        add_bookmark(doc.paragraphs[idx], name, bid)
        heading_for_toc[toc_label] = name
        bid += 1

    matched = 0
    for row in toc.rows[1:]:
        # Top-level rows carry their number in the first column ("1." +
        # "INTRODUCTION"), sub-rows carry the whole label in the second.
        num = row.cells[0].text.strip()
        label = row.cells[1].text.strip()
        if num:
            label = f"{num} {label}".strip()
        if label in heading_for_toc:
            set_pageref(row.cells[2], heading_for_toc[label])
            matched += 1
        else:
            row.cells[2].text = ""
    print(f"[13] Table of Contents: labels corrected, {matched} PAGEREF "
          f"fields inserted")

    # =======================================================================
    # 14. Final structural formatting: each numbered chapter starts on its own
    #     page, and the front matter is numbered separately from the body.
    # =======================================================================
    for chapter in ("ABSTRACT", "TABLE OF CONTENTS", "1. INTRODUCTION",
                    "2. PROJECT DESCRIPTION AND GOALS",
                    "3. TECHNICAL SPECIFICATION",
                    "4. DESIGN APPROACH AND DETAILS", "5. REFERENCES"):
        page_break_before(doc.paragraphs[find(doc, chapter)])
    set_roman_front_matter(doc)

    # Headings stay attached to the text (or figure) that follows them.
    for p in doc.paragraphs:
        t = p.text.strip()
        runs = [r for r in p.runs if r.text.strip()]
        is_heading = bool(runs) and runs[0].bold and runs[0].font.size is not None
        if is_heading or re.match(r"^\d\.\d(\.\d)? ", t):
            keep_with_next(p)
    print("[14] chapter page breaks, keep-with-next headings, roman front matter")

    n = normalise_body_text(doc)
    for obj in ("Parameter Schema:", "Local Language Understanding:",
                "Validation:", "Pipeline Demonstration:",
                "Future Extensibility:"):
        try:
            bold_lead_in(doc.paragraphs[find(doc, obj)], obj[:-1])
        except LookupError:
            print(f"    ! objective not found: {obj}")

    b = normalise_bullets(doc)
    print(f"[14b] normalised {n} body paragraphs and {b} bulleted items "
          f"(justified, bold labels, 1.15 spacing)")

    # Collapse runs of consecutive empty paragraphs left behind by the edits.
    blanks = 0
    prev_blank = False
    cover_end = find(doc, "ABSTRACT")
    for p in list(doc.paragraphs[cover_end:]):
        is_blank = (p.text.strip() == ""
                    and not p._p.findall(".//" + qn("a:blip")))
        if is_blank and prev_blank:
            delete_paragraph(p)
            blanks += 1
        else:
            prev_blank = is_blank
    print(f"[15] collapsed {blanks} duplicated blank paragraphs")

    # =======================================================================
    doc.save(OUT)
    print(f"\nSaved: {OUT}")


if __name__ == "__main__":
    main()
