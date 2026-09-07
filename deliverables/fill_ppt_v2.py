"""
Fills the Review-2 PPT template with the APPROVED, reframed content
(SLM / natural-language vision primary, CPU/C++ explicitly labeled as
this review's prototype, CUDA/GPU as the intended direction). Builds on
the same template-preserving approach as fill_ppt.py: edit existing runs
in place, clone paragraphs (or, for the literature review, insert a real
table) rather than inventing new formatting.
"""
import copy
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

SRC = r"C:\Users\SHIVANI\Downloads\4 Review 2 PPT Template (1).pptx"
OUT = os.path.join(os.path.dirname(__file__), "Review2_PPT_Final.pptx")
FIG = os.path.join(os.path.dirname(__file__), "figures")

prs = Presentation(SRC)
slide_list = list(prs.slides)


def slide(n):
    return slide_list[n - 1]


def shape_by_ph_idx(sl, name_contains="Content Placeholder"):
    for shp in sl.shapes:
        if shp.has_text_frame and name_contains in shp.name:
            return shp
    return None


def shape_by_name(sl, name):
    for shp in sl.shapes:
        if shp.name == name:
            return shp
    return None


def set_run_text(paragraph, text):
    runs = paragraph.runs
    if not runs:
        paragraph.text = text
        return
    runs[0].text = text
    for r in runs[1:]:
        r.text = ""


def clone_paragraph_after(paragraph, text):
    new_p = copy.deepcopy(paragraph._p)
    paragraph._p.addnext(new_p)
    from pptx.text.text import _Paragraph
    new_para = _Paragraph(new_p, paragraph._parent)
    set_run_text(new_para, text)
    return new_para


def fill_bullets(shape, lines):
    tf = shape.text_frame
    template_paras = list(tf.paragraphs)
    n_existing = len(template_paras)
    for i, line in enumerate(lines):
        if i < n_existing:
            set_run_text(template_paras[i], line)
        else:
            template_paras.append(clone_paragraph_after(template_paras[-1], line))
    for j in range(len(lines), n_existing):
        set_run_text(template_paras[j], "")


# ============================================================
# SLIDE 1 -- Title (unchanged content, same as before)
# ============================================================
s = slide(1)
title_shape = None
team_shape = None
for shp in s.shapes:
    if not shp.has_text_frame:
        continue
    if "<<Project Title>>" in shp.text_frame.text:
        title_shape = shp
    elif "Team" in shp.text_frame.text and "Name" in shp.text_frame.text:
        team_shape = shp

for p in title_shape.text_frame.paragraphs:
    for r in p.runs:
        if r.text == "<<Programme ":
            r.text = "B.Tech (CSE)"
        elif r.text in ("B.Tech", "M.Tech", "/ "):
            r.text = ""
        elif "Course Code" in r.text:
            r.text = "  |  BCSE497J - Project-I"
        elif r.text == "<<Project Title>>":
            r.text = ("AI-Assisted Generation of CUDA-Based Molecular Dynamics "
                       "Simulations from Natural-Language Descriptions")

name_done = False
for p in team_shape.text_frame.paragraphs:
    for r in p.runs:
        if r.text.startswith("<<Name >>") and not name_done:
            r.text = "Shivani Srivastava\t\t  ("
        elif r.text == "Reg.No" and not name_done:
            r.text = "23BDS0085"
        elif r.text == ">>)" and not name_done:
            r.text = ")"
            name_done = True
        elif r.text == "<<":
            r.text = ""
        elif r.text == "Guidename":
            r.text = "Dr. Vijayasherly V"
        elif r.text == ">>":
            r.text = ""

# ============================================================
# SLIDE 3 -- Aim
# ============================================================
s = slide(3)
shp = shape_by_ph_idx(s)
set_run_text(shp.text_frame.paragraphs[0],
    "To develop an AI-assisted molecular dynamics simulation framework that allows users "
    "to describe simulation requirements in natural language, automatically interprets and "
    "configures those requirements, and reduces the effort required to build and run "
    "molecular dynamics simulations \u2014 enabling users to focus more on scientific research "
    "and understanding rather than manually writing simulation code and configurations.")

# ============================================================
# SLIDE 4 -- Abstract
# ============================================================
s = slide(4)
shp = shape_by_ph_idx(s)
fill_bullets(shp, [
    "Setting up a molecular dynamics (MD) simulation traditionally demands both deep "
    "physics knowledge and low-level, high-performance programming expertise \u2014 a "
    "combination that puts GPU-accelerated MD out of reach for many students and "
    "researchers who understand the science but not the systems programming underneath "
    "it. This project proposes an AI-assisted MD framework in which a local Small "
    "Language Model (SLM) reads a plain-English description of a desired simulation, "
    "extracts the physical parameters and settings it implies, and automatically "
    "configures a simulation accordingly \u2014 letting the user focus on the scientific "
    "question rather than the mechanics of setting one up.",
    "The language component's role is deliberately restricted to understanding intent "
    "and extracting structured parameters; it never generates or modifies simulation "
    "physics itself, which is instead handled by a fixed, hand-verified simulation "
    "engine. The intended computational direction for this framework is GPU/CUDA "
    "acceleration. As a working demonstration, a CPU-based C++ prototype has been "
    "implemented \u2014 supporting Lennard-Jones and Hard-Sphere models, multiple boundary "
    "conditions, and solid/liquid/gas states \u2014 with real trajectories shown through an "
    "interactive 3D visualization, validating the pipeline end to end.",
])

# ============================================================
# SLIDE 5 -- Literature Review (real table, 12 papers)
# ============================================================
s = slide(5)
LIT_ROWS = [
    ("Paper / Author", "Objective", "Methodology", "Pros / Cons", "Key Relevance"),
    ("Frenkel & Smit (2002)", "Establish classical MD algorithms",
     "Lennard-Jones potential, Verlet integrators, PBC",
     "Pro: foundational. Con: no automation",
     "Physical basis for this project's engine"),
    ("Plimpton, J. Comput. Phys. (1995) \u2014 LAMMPS", "Scalable parallel MD",
     "Spatial domain decomposition",
     "Pro: mature, scalable. Con: no NL interface",
     "Motivates a simpler NL entry point"),
    ("Abraham et al., SoftwareX (2015) \u2014 GROMACS", "High-performance biomolecular MD",
     "Multi-level CPU+GPU parallelism",
     "Pro: highly optimized. Con: manual scripting",
     "Raw performance alone doesn't solve configuration effort"),
    ("Stone et al., J. Comput. Chem. (2007)", "Early GPU acceleration of MD",
     "CUDA force/energy kernels",
     "Pro: large GPU speedups. Con: deep CUDA expertise needed",
     "Direct precedent for this project's CUDA direction"),
    ("Anderson et al., J. Comput. Phys. (2008) \u2014 HOOMD-blue", "Fully GPU-resident MD engine",
     "Custom CUDA kernels for whole loop",
     "Pro: no CPU bottleneck. Con: complex internals",
     "Reference architecture for a CUDA port"),
    ("Eastman et al., PLOS Comput. Biol. (2017) \u2014 OpenMM 7", "Flexible GPU-accelerated MD toolkit",
     "Modular plugins over CUDA/OpenCL",
     "Pro: flexible, adopted. Con: still API-driven",
     "Even flexible toolkits stop at an API, not language"),
    ("Behler & Parrinello, PRL (2007)", "Learn interatomic potentials via neural nets",
     "NN-fitted potential energy surfaces",
     "Pro: near-quantum accuracy. Con: needs large datasets",
     "Contrast: ML inside physics vs. this project's ML outside physics"),
    ("Butler et al., Nature (2018)", "Survey ML across molecular/materials science",
     "Literature review",
     "Pro: broad, well-cited. Con: not automation-specific",
     "General context for ML's role in this space"),
    ("No\u00e9 et al., Annu. Rev. Phys. Chem. (2020)", "Survey ML methods for MD",
     "Review of ML-for-MD techniques",
     "Pro: comprehensive. Con: physics-internal ML only",
     "Confirms NL interfaces for MD remain unexplored"),
    ("White et al., arXiv:2302.11382 (2023)", "Catalog reusable prompt-design patterns",
     "Pattern-based prompt engineering",
     "Pro: improves reliability. Con: needs downstream validation",
     "Basis for constraining language output to a fixed schema"),
    ("M. Bran et al., Nat. Mach. Intell. (2024) \u2014 ChemCrow", "LLM agent calling verified chemistry tools",
     "LLM + tool-calling architecture",
     "Pro: keeps LLM out of computation. Con: limited by tool coverage",
     "Mirrors this project's core design: extraction vs. verified execution"),
    ("Boiko et al., Nature (2023) \u2014 Coscientist", "LLM-driven planning of real experiments",
     "LLM orchestrating external tools",
     "Pro: real tasks via NL. Con: reliability depends on tools called",
     "Evidence NL-driven scientific workflows are a credible direction"),
]

content_shape = shape_by_name(s, "Content Placeholder 4")
left = content_shape.left
top = content_shape.top
width = content_shape.width
height = content_shape.height
n_rows, n_cols = len(LIT_ROWS), 5
graphic_frame = s.shapes.add_table(n_rows, n_cols, left, top, width, height)
table = graphic_frame.table
col_weights = [0.20, 0.19, 0.18, 0.20, 0.23]
for i, w in enumerate(col_weights):
    table.columns[i].width = int(width * w)

for r, row_vals in enumerate(LIT_ROWS):
    for c, val in enumerate(row_vals):
        cell = table.cell(r, c)
        cell.text = val
        for p in cell.text_frame.paragraphs:
            p.font.size = Pt(8) if r > 0 else Pt(9)
            p.font.name = "Times New Roman"
            p.font.bold = (r == 0)
            if r == 0:
                p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(0x2F, 0x6F, 0x5E) if r == 0 else (
            RGBColor(0xF3, 0xF7, 0xF5) if r % 2 == 0 else RGBColor(0xFF, 0xFF, 0xFF))
        cell.margin_top = Pt(2)
        cell.margin_bottom = Pt(2)
        cell.margin_left = Pt(3)
        cell.margin_right = Pt(3)
table.rows[0].height = Inches(0.35)
for r in range(1, n_rows):
    table.rows[r].height = Inches((height / 914400 - 0.35) / (n_rows - 1) * 914400 / 914400 * 1.0)

# hide the now-unused empty content placeholder behind the table
content_shape._element.getparent().remove(content_shape._element)

# ============================================================
# SLIDE 6 -- Research Gap
# ============================================================
s = slide(6)
shp = shape_by_ph_idx(s)
fill_bullets(shp, [
    "GPU-accelerated MD toolkits (LAMMPS, GROMACS, OpenMM, HOOMD-blue) are powerful but "
    "assume the user already knows how to configure a simulation, or write CUDA for "
    "custom behavior \u2014 no natural-language entry point exists anywhere in this stack.",
    "General-purpose LLMs can generate code, but nothing constrains the generated code "
    "to be physically correct \u2014 unconstrained LLM-generated low-level/parallel code is "
    "known to be unreliable.",
    "Emerging LLM-tool-use systems (ChemCrow, Coscientist) show natural language can "
    "drive real scientific workflows when the LLM is restricted to calling verified "
    "tools \u2014 but this pattern hasn't been applied to GPU-accelerated MD specifically.",
    "This project addresses that gap: a local language-understanding component "
    "restricted to parameter extraction only, paired with a fixed, hand-verified "
    "simulation engine, with GPU/CUDA acceleration as the intended target.",
])

# ============================================================
# SLIDE 7 -- Objectives (+ SDG/Outcomes line)
# ============================================================
s = slide(7)
shp = shape_by_ph_idx(s)
fill_bullets(shp, [
    "Design a structured parameter schema capturing a full MD simulation request \u2014 "
    "particle count, density/temperature, interaction model, state of matter, "
    "boundary/container settings.",
    "Build a local language-understanding component converting natural language into "
    "this schema, with no external AI API \u2014 validated here with a rule-based extractor, "
    "a full local SLM identified as the target upgrade.",
    "Implement a validation layer checking physical feasibility before any simulation "
    "runs.",
    "Demonstrate the complete natural-language \u2192 configured \u2192 executed \u2192 visualized "
    "pipeline on a CPU-based prototype, proving the pipeline design is sound.",
    "Establish a simulation engine architecture whose core computation is structured for "
    "a future CUDA/GPU port, carrying the accessibility gains to large-scale simulation.",
])
for shp2 in s.shapes:
    if shp2.has_text_frame and "Identified" in shp2.text_frame.text:
        for p in shp2.text_frame.paragraphs:
            for r in p.runs:
                if r.text.startswith("Identified:"):
                    r.text = ("Identified:  Sustainable Development Goals \u2014 SDG 9 (Primary), "
                               "SDG 4 (Secondary)   |   ")
                elif r.text == "Outcomes (":
                    r.text = "Outcomes: "
                elif "(Patent / Conference" in r.text:
                    r.text = "Product (Working Prototype / Proof of Concept)"
                elif "Journal (Scopus)" in r.text:
                    r.text = ""

# ============================================================
# SLIDE 8 -- Framework / Architecture (backend-agnostic pipeline)
# ============================================================
s = slide(8)
shp = shape_by_ph_idx(s)
fill_bullets(shp, [
    "The pipeline is designed so the execution backend is swappable: the CPU engine "
    "built for this review demonstrates the full pipeline end-to-end, while the same "
    "configuration format is intended to drive a CUDA-accelerated engine for large-scale "
    "runs. The language-understanding step's role is fixed regardless of backend \u2014 it "
    "only ever produces parameters, never simulation code.",
])
s.shapes.add_picture(os.path.join(FIG, "fig_pipeline.png"),
                      left=Inches(0.7), top=Inches(2.7), width=Inches(8.4))

# ============================================================
# SLIDE 9 -- Functional Requirements
# ============================================================
s = slide(9)
shp = shape_by_ph_idx(s)
fill_bullets(shp, [
    "Accept a plain-English simulation description and extract structured parameters "
    "locally, with no external AI API.",
    "Accept manual parameter entry as an equivalent alternative to natural-language "
    "input.",
    "Validate all parameters, from either input path, against physical feasibility "
    "before configuring a simulation.",
    "Translate validated parameters into a complete simulation configuration "
    "automatically, without the user writing simulation code.",
    "Prototype: execute the configured simulation on CPU (Lennard-Jones / Hard-Sphere, "
    "periodic or container boundaries).",
    "Record real trajectories and energy, and present them through an interactive 3D "
    "visualization.",
    "Future scope: keep the configuration format backend-agnostic, so it can drive a "
    "CUDA/GPU engine without redesign.",
])

# ============================================================
# SLIDE 10 -- Modules
# ============================================================
s = slide(10)
shp = shape_by_ph_idx(s)
fill_bullets(shp, [
    "Natural-Language Understanding Module \u2014 local, offline; interprets intent. "
    "(Prototype: rule-based. Target: local SLM.)",
    "Parameter Validation Module \u2014 checks physical feasibility for either input path.",
    "Simulation Configuration Module \u2014 turns validated parameters into a ready-to-run "
    "setup; where \u201cthe user stops writing code\u201d happens.",
    "Simulation Execution Module \u2014 runs the configuration. (Prototype: CPU/C++, this "
    "review. Target: CUDA/GPU.)",
    "Output & Visualization Module \u2014 captures real trajectory/energy data and renders "
    "an interactive 3D view.",
])

# ============================================================
# SLIDE 11 -- Experiments and Results
# ============================================================
s = slide(11)
shp = shape_by_ph_idx(s)
fill_bullets(shp, [
    "Full pipeline demonstrated end-to-end: description \u2192 extracted parameters \u2192 "
    "validated configuration \u2192 executed CPU simulation \u2192 real 3D visualization, no "
    "manual coding step.",
    "21 automated tests confirm parameter extraction and validation behave correctly "
    "across both manual and natural-language input paths.",
    "CPU prototype physics verified correct \u2014 periodic (NVE) runs conserve total energy "
    "to within 0.5 reduced units; Hard-Sphere collisions conserve kinetic energy almost "
    "exactly.",
    "Solid, liquid, and gas requests produce visibly distinct, physically appropriate "
    "behavior \u2014 confirming NL state descriptions map onto real configuration, not "
    "cosmetic presets.",
    "Not yet demonstrated: GPU/CUDA execution (hardware not available) and a full local "
    "SLM (prototype uses a rule-based stand-in for this pipeline stage).",
])

# ============================================================
# SLIDE 12 -- Conclusion
# ============================================================
s = slide(12)
shp = shape_by_ph_idx(s)
fill_bullets(shp, [
    "This project demonstrates that a natural-language, AI-assisted front end to MD "
    "simulation is achievable without compromising physical correctness \u2014 by strictly "
    "separating \u201cunderstanding intent\u201d from \u201ccomputing the physics.\u201d",
    "A working CPU-based prototype validates this pipeline end-to-end: description or "
    "manual input \u2192 automatic configuration \u2192 real simulation \u2192 real 3D results.",
    "The broader goal \u2014 accessible, GPU-accelerated MD without deep CUDA expertise \u2014 "
    "remains the project's intended direction, with this prototype establishing the "
    "architecture it will be built on.",
])

# ============================================================
# SLIDE 13 -- References (13 entries)
# ============================================================
s = slide(13)
REFERENCES = [
    "D. Frenkel and B. Smit, Understanding Molecular Simulation: From Algorithms to "
    "Applications, 2nd ed. San Diego, CA: Academic Press, 2002.",
    "S. Plimpton, \u201cFast parallel algorithms for short-range molecular dynamics,\u201d "
    "J. Comput. Phys., vol. 117, no. 1, pp. 1-19, 1995.",
    "M. J. Abraham et al., \u201cGROMACS: High performance molecular simulations through "
    "multi-level parallelism from laptops to supercomputers,\u201d SoftwareX, vol. 1-2, "
    "pp. 19-25, 2015.",
    "J. E. Stone et al., \u201cAccelerating molecular modeling applications with graphics "
    "processors,\u201d J. Comput. Chem., vol. 28, no. 16, pp. 2618-2640, 2007.",
    "J. A. Anderson, C. D. Lorenz, and A. Travesset, \u201cGeneral purpose molecular dynamics "
    "simulations fully implemented on graphics processing units,\u201d J. Comput. Phys., "
    "vol. 227, no. 10, pp. 5342-5359, 2008.",
    "P. Eastman et al., \u201cOpenMM 7: Rapid development of high performance algorithms for "
    "molecular dynamics,\u201d PLOS Comput. Biol., vol. 13, no. 7, 2017.",
    "J. Behler and M. Parrinello, \u201cGeneralized neural-network representation of "
    "high-dimensional potential-energy surfaces,\u201d Phys. Rev. Lett., vol. 98, 146401, 2007.",
    "K. T. Butler, D. W. Davies, H. Cartwright, O. Isayev, and A. Walsh, \u201cMachine "
    "learning for molecular and materials science,\u201d Nature, vol. 559, pp. 547-555, 2018.",
    "F. No\u00e9, A. Tkatchenko, K.-R. M\u00fcller, and C. Clementi, \u201cMachine learning for "
    "molecular simulation,\u201d Annu. Rev. Phys. Chem., vol. 71, pp. 361-390, 2020.",
    "J. White et al., \u201cA prompt pattern catalog to enhance prompt engineering with "
    "ChatGPT,\u201d arXiv:2302.11382, 2023.",
    "A. M. Bran, S. Cox, O. Schilter, C. Baldassari, A. D. White, and P. Schwaller, "
    "\u201cAugmenting large language models with chemistry tools,\u201d Nature Machine "
    "Intelligence, 2024.",
    "D. A. Boiko, R. MacKnight, B. Kline, and G. Gomes, \u201cAutonomous chemical research "
    "with large language models,\u201d Nature, vol. 624, pp. 570-578, 2023.",
    "NVIDIA Corporation, CUDA C++ Programming Guide, Release 12.5, NVIDIA, 2024.",
]
ref_shape = None
for shp in s.shapes:
    if shp.has_text_frame and shp.text_frame.text.strip().startswith("[1]"):
        ref_shape = shp
        break
# The template's 2nd reference ("[2] A. B. Researcher ... Proc. Name of
# Conference ... City, Country ...") is a separate, still-untouched
# paragraph further down this shape -- blank it now, before cloning, so
# it doesn't linger as a leftover after the 13 real references.
for p in ref_shape.text_frame.paragraphs:
    if "Proc. Name of Conference" in p.text or p.text.strip().startswith("[2]"):
        set_run_text(p, "")

p0 = ref_shape.text_frame.paragraphs[0]
set_run_text(p0, "[1]  " + REFERENCES[0])
last = p0
for i, ref in enumerate(REFERENCES[1:], start=2):
    last = clone_paragraph_after(last, f"[{i}]  {ref}")

prs.save(OUT)
print("Saved:", OUT)
