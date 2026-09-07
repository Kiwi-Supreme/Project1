"""
Final report build: reframed around the SLM/CUDA vision, per the approved
content proposal, plus three requested edits:
  1. Motivation section tells the actual internship story.
  2. Hardware Specification lists GPU as a real requirement for the
     full/final framework (not hedged as "not required").
  3. Design section includes all four diagrams (DFD, Use Case, Class,
     Sequence) redrawn/relabeled at the framework level.

Same template-preserving approach as before: edit existing runs in place,
clone paragraphs for anything needing more lines than the template had.
"""
import copy
import os
from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

SRC = r"C:\Users\SHIVANI\Downloads\3 BCSE497J Project I Report - Template (1).docx"
OUT = os.path.join(os.path.dirname(__file__), "BCSE497J_Project_I_Report_Final.docx")
FIG = os.path.join(os.path.dirname(__file__), "figures")

doc = Document(SRC)


def collect_paragraphs(document):
    result = list(document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                result.extend(cell.paragraphs)
    return result


def set_text(paragraph, text):
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(text)
        return
    runs[0].text = text
    for r in runs[1:]:
        r.text = ""


def find_para(paragraphs, prefix, start=0):
    for i in range(start, len(paragraphs)):
        if paragraphs[i].text.startswith(prefix):
            return i
    raise ValueError("paragraph not found: " + prefix)


def clone_after(paragraph, text=None):
    new_p_elem = copy.deepcopy(paragraph._p)
    ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    for tag in ("bookmarkStart", "bookmarkEnd"):
        for el in new_p_elem.findall(ns + tag):
            new_p_elem.remove(el)
    paragraph._p.addnext(new_p_elem)
    from docx.text.paragraph import Paragraph
    new_para = Paragraph(new_p_elem, paragraph._parent)
    if text is not None:
        set_text(new_para, text)
    return new_para


def clone_many_after(paragraph, texts):
    anchor = paragraph
    for t in texts:
        anchor = clone_after(anchor, t)
    return anchor


def add_picture_after(paragraph, image_path, width_in=5.6):
    new_p = clone_after(paragraph, "")
    new_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    new_p.add_run().add_picture(image_path, width=Inches(width_in))
    return new_p


paras = collect_paragraphs(doc)

# ============================================================
# COVER PAGE
# ============================================================
set_text(paras[find_para(paras, "<Title of Project>")],
         "AI-ASSISTED GENERATION OF CUDA-BASED MOLECULAR DYNAMICS "
         "SIMULATIONS FROM NATURAL-LANGUAGE DESCRIPTIONS")
set_text(paras[find_para(paras, "Reg. No. 1")], "23BDS0085")
set_text(paras[find_para(paras, "Student Name 1")], "SHIVANI SRIVASTAVA")
set_text(paras[find_para(paras, "Project Guide Name")], "Dr. Vijayasherly V")

# ============================================================
# ABSTRACT
# ============================================================
i = find_para(paras, "The rapid advancement of technology")
set_text(paras[i],
    "Setting up a molecular dynamics (MD) simulation traditionally demands both deep "
    "computational-physics knowledge and low-level, high-performance programming expertise "
    "\u2014 a combination that puts GPU-accelerated MD out of reach for many students and "
    "researchers who understand the science but not the systems programming underneath it. "
    "This project proposes an AI-assisted molecular dynamics framework in which a local "
    "Small Language Model (SLM) reads a plain-English description of a desired simulation, "
    "extracts the physical parameters and settings it implies, and automatically configures "
    "a simulation accordingly \u2014 letting a researcher focus on the scientific question "
    "rather than the mechanics of setting one up.")
i2 = find_para(paras, "ML algorithms, such as supervised")
set_text(paras[i2],
    "The language-understanding component's role is deliberately restricted to "
    "interpreting intent and extracting structured parameters; it never generates or "
    "modifies simulation physics itself, which remains the responsibility of a fixed, "
    "independently verifiable simulation engine. This separation is the framework's central "
    "design principle: natural language lowers the barrier to entry without lowering trust "
    "in the result. The proposed computational direction for the framework is GPU/CUDA "
    "acceleration, so that the accessibility gained through natural-language configuration "
    "extends to large-scale, high-performance simulation.")
i3 = find_para(paras, "Through case studies and experimental")
set_text(paras[i3],
    "As a working demonstration of this concept, a CPU-based prototype has been "
    "implemented, validating the complete pipeline \u2014 natural-language or manual input, "
    "automatic parameter extraction and validation, simulation configuration, execution, and "
    "interactive 3D visualization of the real resulting trajectory \u2014 end to end. The "
    "prototype supports Lennard-Jones and Hard-Sphere interaction models, multiple boundary "
    "conditions, and physically distinct solid/liquid/gas states, and establishes the "
    "architecture on which the full, GPU-accelerated framework is intended to be built.")
i4 = find_para(paras, "Keywords - Artificial Intelligence")
set_text(paras[i4],
    "Keywords - Molecular Dynamics, Small Language Model, Natural-Language Interface, "
    "Automated Simulation Configuration, GPU/CUDA Acceleration, Accessible Scientific "
    "Computing.")

# ============================================================
# 1.1 BACKGROUND
# ============================================================
i = find_para(paras, "The rapid digitization of various sectors")
set_text(paras[i],
    "Molecular dynamics (MD) simulation is a well-established computational technique for "
    "studying how atoms and particles move and interact over time, with applications "
    "spanning materials science, drug design, and condensed-matter physics. Running a "
    "meaningful MD simulation \u2014 particularly a GPU-accelerated one \u2014 requires a "
    "researcher to translate a scientific question into simulation code and configuration: "
    "choosing an interaction model, initializing a system correctly, setting boundary "
    "conditions, and writing or configuring numerical integration, often in a low-level, "
    "high-performance language such as C++ or CUDA.")
i = find_para(paras, "In response to these challenges, Machine Learning")
set_text(paras[i],
    "This technical overhead is largely independent of the scientific question being "
    "asked. Two researchers investigating very different physical systems still have to "
    "solve the same configuration and implementation problems before either can begin "
    "producing results. Mature MD packages such as LAMMPS and GROMACS reduce some of this "
    "burden through configuration scripting, but scripting remains a technical skill "
    "distinct from the physics itself, and neither offers a natural-language entry point.")
i = find_para(paras, "This project seeks to explore the role of ML")
set_text(paras[i],
    "This project explores whether a locally-run, language-understanding component can "
    "remove that overhead: reading a plain-English description of a desired simulation, "
    "extracting the parameters it implies, and configuring a simulation automatically \u2014 "
    "so that the researcher's time and attention go to the scientific question, not the "
    "setup work surrounding it.")

# ============================================================
# 1.2 MOTIVATION -- the internship story
# ============================================================
i = find_para(paras, "The ever-evolving landscape of cyber threats")
set_text(paras[i],
    "The motivation for this project comes from direct, personal experience. During a "
    "summer internship in the Biotech Department at IIT Bombay (25th May \u2013 11th July 2024), "
    "under the guidance of Prof. Ranjith Padinhateeri, I hand-wrote and debugged a "
    "Lennard-Jones molecular dynamics simulation in C++, and then attempted to port that "
    "same simulation to CUDA to take advantage of GPU acceleration. The C++ implementation "
    "worked correctly, reproducing the expected energy-conservation behaviour of an NVE "
    "simulation. The CUDA port, however, ran into real, unresolved problems \u2014 an incorrect "
    "memory-allocation cast, unsupported double-precision atomic operations on the target "
    "GPU architecture, and incomplete host/device memory management \u2014 that I was ultimately "
    "unable to fully resolve within the internship's timeframe, given my limited CUDA "
    "experience at the time.")
i = find_para(paras, "This project is motivated by the potential")
set_text(paras[i],
    "That experience directly motivated this project. It made two things clear. First, "
    "that a large share of the effort in building a GPU-accelerated MD simulation lies not "
    "in the physics itself \u2014 which is comparatively well-established and reusable \u2014 but in "
    "the systems-programming work of correctly wiring together initialization, force "
    "calculation, memory management, and numerical integration, particularly on the GPU. "
    "Second, that this repetitive, pattern-following configuration work is exactly the kind "
    "of task a language-understanding system could reasonably absorb, provided it is kept "
    "strictly separate from the parts of the problem \u2014 the physics, and the CUDA "
    "implementation itself \u2014 that genuinely require verified correctness, the way my own "
    "CUDA debugging did.")
i = find_para(paras, "The goal of this project is to explore how ML")
set_text(paras[i],
    "This also motivated the choice of a local SLM over a cloud-based AI assistant: it "
    "keeps the system self-contained, explainable, reproducible, and usable without an "
    "internet connection or ongoing external dependency \u2014 properties that matter for a tool "
    "meant to be understood, verified, and demonstrated end to end, in the same spirit as "
    "the hand-verified C++ engine it configures, rather than trusted as a black box.")

# ============================================================
# 1.3 SCOPE OF THE PROJECT
# ============================================================
i = find_para(paras, "The scope of this project focuses on exploring")
set_text(paras[i],
    "The proposed system is a natural-language and manual-input interface, backed by a "
    "local SLM restricted to parameter extraction, driving a GPU/CUDA-accelerated molecular "
    "dynamics engine, with interactive visualization of results. This project's scope covers "
    "a structured parameter schema (particle count, density/temperature or speed, "
    "interaction model, state of matter, boundary/container configuration), a local "
    "natural-language parameter-extraction component and an equivalent manual-entry path, "
    "both checked by one shared feasibility validator.")
i = find_para(paras, "The project will involve the development")
set_text(paras[i],
    "It further covers a simulation engine implementing two interaction models "
    "(Lennard-Jones, Hard-Sphere), two boundary types (periodic, and a real cuboidal/"
    "spherical container), and three states of matter represented as genuine physical "
    "differences rather than cosmetic presets, plus an interactive 3D visualization of the "
    "real simulated trajectory and its energy behaviour over time, as evidence the "
    "simulation is physically correct.")
i = find_para(paras, "The project also considers the ethical")
set_text(paras[i],
    "GPU/CUDA execution is the framework's intended computational direction but is not "
    "implemented at this project stage \u2014 demonstrated instead via a CPU-based prototype "
    "covering the same pipeline, given the hardware available for development and "
    "demonstration. A full local SLM is the framework's intended language component; this "
    "stage validates the same pipeline position with a lighter, rule-based extractor. Both "
    "are future extensions of this same architecture, not separate unfinished features.")

# ============================================================
# 2.1 LITERATURE REVIEW -- full academic prose, 12 sources
# ============================================================
anchor = paras[find_para(paras, "< Maximum of 50 journal papers")]
clone_many_after(anchor, [
    "Foundations of molecular dynamics. Frenkel and Smit\u2019s Understanding Molecular "
    "Simulation [1] remains the standard reference for the algorithms this project's "
    "simulation engine is built on: the Lennard-Jones potential, Verlet-family integrators, "
    "and periodic boundary conditions. These form the physical and numerical basis for any "
    "MD engine, including the one implemented here, and ground the project's parameter "
    "schema in established, well-understood physics rather than an ad hoc simplification.",
    "Performance and GPU/CUDA acceleration. Large-scale MD packages demonstrate both the "
    "payoff and the cost of high-performance MD. Plimpton\u2019s LAMMPS [2] established spatial "
    "domain decomposition for parallel MD at scale; Abraham et al.\u2019s GROMACS [3] extended "
    "this with multi-level CPU/GPU parallelism for biomolecular systems. Both remain "
    "configuration-driven, script-based tools with no natural-language interface. On the "
    "GPU side specifically, Stone et al. [4] demonstrated early, substantial GPU speedups "
    "for MD force calculations using CUDA, and Anderson, Lorenz, and Travesset\u2019s HOOMD-blue "
    "[5] pushed further, implementing an entire MD simulation loop natively on the GPU. "
    "Eastman et al.\u2019s OpenMM [6] shows that even a deliberately flexible, plugin-based GPU "
    "MD toolkit still stops at a programmatic API \u2014 confirming that raw computational "
    "performance and ease of use are separate problems, and that the second remains largely "
    "unsolved even in mature, high-performance tools. Collectively, this body of work both "
    "motivates GPU/CUDA as this project's intended computational direction and confirms "
    "that no existing high-performance MD tool addresses the configuration-accessibility "
    "problem this project targets.",
    "Machine learning inside molecular simulation. A separate line of research applies "
    "machine learning within the physics itself \u2014 Behler and Parrinello\u2019s neural-network "
    "potential energy surfaces [7] are the seminal example, learning interatomic potentials "
    "that approach quantum-mechanical accuracy at MD-compatible speed. Butler et al. [8] and "
    "No\u00e9 et al. [9] survey this broader intersection of machine learning and molecular/"
    "materials science. This project takes a deliberately different position: rather than "
    "using learned models to approximate physics, it keeps machine learning entirely "
    "outside the physics, restricted to interpreting a user's intent. Reviewing this "
    "literature clarifies that distinction rather than undermining it \u2014 the surveys in "
    "[8, 9] confirm that natural-language interfaces to simulation setup remain "
    "comparatively unexplored relative to ML-for-physics work.",
    "Language models, structured extraction, and tool-restricted scientific agents. White "
    "et al.\u2019s prompt-pattern catalog [10] and related structured-output/function-calling "
    "techniques underpin the decision to constrain this project's language component to a "
    "fixed, validated parameter schema rather than free-form generation. Most directly "
    "relevant are two recent systems that apply this \u201clanguage model as intent-interpreter, "
    "verified tool as executor\u201d pattern to real science: ChemCrow [11], where an LLM agent "
    "calls verified chemistry tools rather than generating chemistry itself, and Coscientist "
    "[12], where an LLM plans and orchestrates real experimental/computational workflows "
    "through external tools. Both demonstrate, in a different domain, the same architectural "
    "principle this project applies to molecular dynamics \u2014 natural language for intent, "
    "verified computation for execution \u2014 and support the claim that this is a credible, "
    "active research direction rather than a novel but untested idea.",
])
paras = collect_paragraphs(doc)

# ============================================================
# 2.2 RESEARCH GAP
# ============================================================
anchor = paras[find_para(paras, "<To identify the areas or aspects")]
clone_after(anchor,
    "Mature, high-performance MD tools (LAMMPS, GROMACS, OpenMM, HOOMD-blue) solve the "
    "performance problem but not the accessibility problem \u2014 all remain script- or "
    "API-configured, with no natural-language entry point. General-purpose large language "
    "models can generate simulation code directly, but nothing constrains that generated "
    "code to be physically correct; unconstrained LLM-generated low-level, parallel code is "
    "documented to be unreliable (race conditions, incorrect memory semantics, silent "
    "numerical errors). Recent LLM-tool-use systems in adjacent scientific domains "
    "(ChemCrow, Coscientist) demonstrate that restricting a language model to "
    "intent-interpretation, backed by verified tools, is a workable and trustworthy pattern "
    "\u2014 but this pattern has not been applied to GPU-accelerated molecular dynamics "
    "specifically. This project addresses that gap directly: a language-understanding "
    "component restricted to parameter extraction only, paired with a fixed, independently "
    "verifiable simulation engine, with GPU/CUDA acceleration as the framework's intended "
    "computational direction.")
paras = collect_paragraphs(doc)

# ============================================================
# 2.3 OBJECTIVES
# ============================================================
anchor = paras[find_para(paras, "< To clearly define what the project")]
clone_many_after(anchor, [
    "Parameter Schema: Design a structured parameter schema capturing a complete molecular "
    "dynamics simulation request \u2014 particle count, density/temperature (or particle "
    "speed), interaction model, state of matter, and boundary/container configuration.",
    "Local Language Understanding: Build a local, natural-language parameter-extraction "
    "component that converts a plain-English simulation description into this schema "
    "without relying on any external AI service.",
    "Validation: Implement a validation layer that checks extracted or manually entered "
    "parameters for physical feasibility before any simulation is configured or run.",
    "Pipeline Demonstration: Demonstrate the complete pipeline \u2014 natural language or "
    "manual input, automatic configuration, execution, and visualization \u2014 end to end on a "
    "working prototype, as evidence the pipeline design itself is sound.",
    "Future Extensibility: Structure the simulation engine's core computation so that it "
    "can be extended to GPU/CUDA acceleration in future work, carrying the accessibility "
    "gains demonstrated here to large-scale, high-performance simulation.",
])
paras = collect_paragraphs(doc)

# ============================================================
# 2.4 PROBLEM STATEMENT
# ============================================================
anchor = paras[find_para(paras, "< Clearly define the issue or challenge")]
clone_after(anchor,
    "Running a molecular dynamics simulation currently requires a researcher to already "
    "possess the systems-programming expertise to translate a physical scenario into "
    "correct simulation code and configuration \u2014 a requirement unrelated to understanding "
    "the underlying physics, and one that excludes many students and early researchers who "
    "understand what they want to simulate but not how to implement it, particularly for "
    "GPU-accelerated simulation. Existing MD tools do not accept a natural-language "
    "description of a desired simulation. This project addresses the problem of providing a "
    "natural-language (and manual) interface to molecular dynamics simulation, without "
    "compromising the physical correctness of the underlying computation by allowing a "
    "language component to generate or influence simulation logic directly \u2014 validating "
    "this approach with a CPU-based prototype ahead of the framework's intended "
    "GPU-accelerated implementation.")
paras = collect_paragraphs(doc)

# ============================================================
# 2.5 PROJECT PLAN (Gantt chart -- reused)
# ============================================================
add_picture_after(paras[find_para(paras, "<A Gantt chart helps")], os.path.join(FIG, "fig1_gantt.png"))

# ============================================================
# 3.1.1 FUNCTIONAL REQUIREMENTS (8 bullets)
# ============================================================
func_items = [
    ("Natural-Language Input", ": The system shall accept a plain-English description of a "
     "desired simulation and extract structured parameters from it using a local "
     "language-understanding component, with no external AI API."),
    ("Manual Parameter Input", ": The system shall accept manually entered parameters as an "
     "equivalent alternative input path."),
    ("Parameter Validation", ": The system shall validate all parameters, from either input "
     "path, against physical feasibility before configuring a simulation, and communicate "
     "specific, human-readable reasons for any rejection."),
    ("Automatic Configuration", ": The system shall translate validated parameters into a "
     "complete, ready-to-run simulation configuration automatically, requiring no manual "
     "simulation coding from the user."),
    ("Simulation Execution", ": The system shall execute the configured simulation and "
     "produce real particle trajectory and energy data. The current prototype executes on "
     "CPU, supporting Lennard-Jones and Hard-Sphere interaction models with periodic or "
     "container (cuboidal/spherical) boundaries."),
    ("State-of-Matter Representation", ": The system shall represent solid, liquid, and gas "
     "states through distinct initial particle arrangement, density, and motion."),
    ("Trajectory and Energy Output", ": The system shall record real particle positions "
     "over time and system energy values for playback and verification."),
    ("Interactive 3D Visualization", ": The system shall present the resulting trajectory "
     "through an interactive 3D visualization, alongside an energy-versus-time view for "
     "verifying physical correctness."),
]
labels = ["Data Collection", "Data Preprocessing", "Model Training", "Real-time Monitoring",
          "Anomaly Detection", "Alert Generation", "User Interface", "Reporting"]
for label, (new_label, new_desc) in zip(labels, func_items):
    idx = find_para(paras, label)
    p = paras[idx]
    p.runs[0].text = new_label
    if len(p.runs) > 1:
        p.runs[1].text = new_desc
        for r in p.runs[2:]:
            r.text = ""
    else:
        p.add_run(new_desc)

# ============================================================
# 3.1.2 NON-FUNCTIONAL REQUIREMENTS (6 bullets)
# ============================================================
nonfunc_items = [
    ("Accessibility", ": The system should not require the user to write or understand "
     "simulation source code to run a simulation."),
    ("Locality and Privacy", ": The system should perform all language understanding and "
     "simulation computation on the local machine, without depending on an external AI API "
     "or network connection."),
    ("Reliability", ": The simulation engine\u2019s physics should remain fixed and "
     "independently verifiable, so its correctness never depends on the reliability of the "
     "language-understanding component."),
    ("Usability", ": The interface should present extracted or entered parameters for "
     "review and confirmation before any simulation runs."),
    ("Maintainability", ": The system should be organized into clearly separated modules "
     "(language understanding, validation, simulation execution, visualization), each "
     "independently understandable and replaceable."),
    ("Extensibility", ": The simulation configuration format and engine should be "
     "structured so the execution backend can be upgraded to GPU/CUDA, and the language "
     "component upgraded to a full local SLM, without redesigning the surrounding system."),
]
nf_labels = ["Performance", "Scalability", "Reliability", "Usability", "Maintainability", "Compliance"]
for label, (new_label, new_desc) in zip(nf_labels, nonfunc_items):
    idx = find_para(paras, label)
    p = paras[idx]
    p.runs[0].text = new_label
    if len(p.runs) > 1:
        p.runs[1].text = new_desc
        for r in p.runs[2:]:
            r.text = ""
    else:
        p.add_run(new_desc)

# ============================================================
# 3.2.1 TECHNICAL FEASIBILITY (4 bullets)
# ============================================================
tf_items = [
    ("Technology Availability", ": The prototype uses well-documented, freely available "
     "technologies (standard C++, Python, plain web frontend technologies), keeping the "
     "current implementation stage low-risk and reproducible on ordinary hardware."),
    ("Technical Expertise", ": The proposed GPU/CUDA-accelerated execution backend is "
     "technically well-precedented but requires NVIDIA GPU hardware not available during "
     "this project stage \u2014 a hardware-driven, not conceptual, limitation."),
    ("Infrastructure", ": The language-understanding component's role (parameter "
     "extraction only) is narrow enough to be reliably validated regardless of whether it "
     "is implemented as a rule-based extractor or a full local SLM, since the simulation "
     "engine independently verifies feasibility either way."),
    ("Integration", ": The frontend, backend, and simulation engine communicate through a "
     "small, well-defined set of local interfaces and configuration formats, keeping "
     "integration straightforward and easy to verify."),
]
tf_labels = ["Technology Availability", "Technical Expertise", "Infrastructure", "Integration"]
for label, (new_label, new_desc) in zip(tf_labels, tf_items):
    idx = find_para(paras, label)
    p = paras[idx]
    p.runs[0].text = new_label
    if len(p.runs) > 1:
        p.runs[1].text = new_desc
        for r in p.runs[2:]:
            r.text = ""
    else:
        p.add_run(new_desc)

# ============================================================
# 3.2.2 ECONOMIC FEASIBILITY (4 bullets)
# ============================================================
ef_items = [
    ("Cost-Benefit Analysis", ": The framework is built entirely on free and open-source "
     "tools; implementation cost is limited to development time, with no recurring "
     "service, licensing, or API cost."),
    ("Budget", ": No paid infrastructure, cloud services, or external AI API subscriptions "
     "are required to build, run, or demonstrate this project."),
    ("Return on Investment (ROI)", ": The primary return is educational and "
     "research-enabling \u2014 a reusable, extensible framework that reduces the setup cost of "
     "future MD studies, rather than a one-off deliverable."),
    ("Funding", ": No external funding is required for this project stage; all development "
     "was carried out using freely available tools."),
]
ef_labels = ["Cost-Benefit Analysis", "Budget", "Return on Investment (ROI)", "Funding"]
for label, (new_label, new_desc) in zip(ef_labels, ef_items):
    idx = find_para(paras, label)
    p = paras[idx]
    p.runs[0].text = new_label
    if len(p.runs) > 1:
        p.runs[1].text = new_desc
        for r in p.runs[2:]:
            r.text = ""
    else:
        p.add_run(new_desc)

# ============================================================
# 3.2.3 SOCIAL FEASIBILITY (4 bullets)
# ============================================================
sf_items = [
    ("User Acceptance", ": Lowering the technical barrier to GPU-accelerated MD is "
     "intended to benefit students and early researchers specifically \u2014 the audience most "
     "limited by systems-programming expertise rather than by scientific understanding."),
    ("Training and Support", ": The project\u2019s documentation records installation, usage, "
     "and the reasoning behind each design decision, so a new user or reviewer can "
     "understand and run the system without additional support."),
    ("Ethical Considerations", ": Running entirely locally, with no data leaving the "
     "user's machine, avoids the privacy and data-governance concerns associated with "
     "cloud-based AI research tools."),
    ("Impact on Workforce", ": The project is an educational and research-enabling tool; "
     "it does not replace or displace any existing workflow or role."),
]
sf_labels = ["User Acceptance", "Training and Support", "Ethical Considerations", "Impact on Workforce"]
for label, (new_label, new_desc) in zip(sf_labels, sf_items):
    idx = find_para(paras, label)
    p = paras[idx]
    p.runs[0].text = new_label
    if len(p.runs) > 1:
        p.runs[1].text = new_desc
        for r in p.runs[2:]:
            r.text = ""
    else:
        p.add_run(new_desc)

# ============================================================
# 3.3.1 HARDWARE SPECIFICATION -- GPU as a real final-project requirement
# ============================================================
hw_map = {
    "Processor": "Processor: A modern multi-core CPU (Intel/AMD, x64) to run the local "
                 "language-understanding component and orchestrate simulation execution.",
    "Memory (RAM)": "Memory (RAM): 8 GB or more recommended, to run language understanding, "
                     "simulation, and visualization concurrently.",
    "Storage": "Storage: Under 500 MB for source code, dependencies, and a local language "
               "model.",
    "Graphics Processing Unit (GPU)": "Graphics Processing Unit (GPU): An NVIDIA GPU with "
        "CUDA support (Compute Capability 6.0 or higher recommended) \u2014 the core "
        "computational requirement for the framework's GPU-accelerated simulation "
        "execution, and the primary hardware this project's architecture is designed "
        "around.",
    "Monitor": "Monitor: A standard display capable of rendering a web browser, for the "
               "interactive 3D visualization.",
}
for label, text in hw_map.items():
    idx = find_para(paras, label)
    set_text(paras[idx], text)

# note distinguishing prototype hardware from the target spec above, placed
# right after the hardware list so it doesn't get lost or misread as a walkback
monitor_para = paras[find_para(paras, "Monitor: A standard display")]
clone_after(monitor_para,
    "Note: the prototype demonstrated for this review was developed and run on CPU-only "
    "hardware; the GPU listed above is the target hardware requirement for the full "
    "framework this project is building toward, not a component of the current "
    "demonstration.")
paras = collect_paragraphs(doc)

# ============================================================
# 3.3.2 SOFTWARE SPECIFICATION
# ============================================================
sw_map = {
    "Operating System:": "Operating System: Windows, Linux, or macOS.",
    "Programming Languages:": "Programming Languages: C++ (simulation engine prototype), "
        "Python (backend/orchestration), HTML/CSS/JavaScript (interface); CUDA C++ for the "
        "proposed GPU-accelerated execution backend.",
    "Development Environment:": "Development Environment: A standard code editor with a "
        "C++ compiler and Python 3.10+ (a CUDA toolkit for the proposed GPU backend).",
    "Libraries and Frameworks:": "Libraries and Frameworks: A lightweight local web "
        "framework for backend orchestration; no external frontend framework is used.",
    "Database:": "Database: None \u2014 the system exchanges structured parameters and "
        "simulation results directly, with no persistent storage layer.",
    "Security Tools:": "Security Tools: Not applicable at this stage \u2014 the system runs "
        "entirely locally with no external network exposure.",
}
for label, text in sw_map.items():
    idx = find_para(paras, label)
    set_text(paras[idx], text)

# ============================================================
# 4.1 SYSTEM ARCHITECTURE -- backend-agnostic pipeline (Fig. 2)
# ============================================================
idx = find_para(paras, "/** Sample **/", start=find_para(paras, "4.1 System Architecture"))
add_picture_after(paras[idx], os.path.join(FIG, "fig_pipeline.png"), width_in=6.2)
paras = collect_paragraphs(doc)
idx = find_para(paras, "<Explanation>")
set_text(paras[idx],
    "Figure 2 shows the proposed system architecture. A user's request \u2014 natural language "
    "or manually entered \u2014 reaches a local language-understanding component whose role is "
    "strictly limited to extracting structured parameters; it never generates or influences "
    "simulation logic. Validated parameters are passed to a fixed simulation-configuration "
    "stage, which is backend-agnostic by design: the same configuration can drive either a "
    "CPU execution engine or a GPU/CUDA-accelerated engine without modification to any "
    "earlier stage. This project stage implements and demonstrates the CPU path; GPU/CUDA "
    "execution is the framework's proposed computational direction. Both paths converge on "
    "the same real trajectory and energy output, visualized identically regardless of which "
    "backend produced it.")

# ============================================================
# 4.2 DESIGN -- all four diagrams
# ============================================================
paras = collect_paragraphs(doc)
idx = find_para(paras, "4.2.1 Data Flow Diagram")
add_picture_after(paras[idx], os.path.join(FIG, "fig3_dfd_v2.png"))
paras = collect_paragraphs(doc)
idx = find_para(paras, "4.2.2 Use Case Diagram")
add_picture_after(paras[idx], os.path.join(FIG, "fig4_usecase.png"))
paras = collect_paragraphs(doc)
idx = find_para(paras, "4.2.3 Class Diagram")
add_picture_after(paras[idx], os.path.join(FIG, "fig5_class_v2.png"))
paras = collect_paragraphs(doc)
idx = find_para(paras, "4.2.4 Sequence Diagram")
add_picture_after(paras[idx], os.path.join(FIG, "fig6_sequence_v2.png"))

# ============================================================
# 5. REFERENCES (13 entries)
# ============================================================
paras = collect_paragraphs(doc)
set_text(paras[find_para(paras, "https://www.cloudflare.com")],
    "https://github.com/Kiwi-Supreme/Molecular-Dynamics-Simulation-of-Particles-in-a-Box")
set_text(paras[find_para(paras, "http://www.thedailybeast.com")],
    "https://platform.openai.com/docs/guides/function-calling")

p = paras[find_para(paras, "Apruzzese, Giovanni")]
p.runs[0].text = "S. Plimpton, \u201cFast parallel algorithms for short-range molecular dynamics,\u201d "
p.runs[1].text = "J. Comput. Phys.,"
p.runs[2].text = " vol. 117, no. 1, pp. 1-19, 1995."
p.runs[4].text = ("M. J. Abraham et al., \u201cGROMACS: High performance molecular simulations "
                   "through multi-level parallelism from laptops to supercomputers,\u201d ")
p.runs[5].text = "SoftwareX,"
p.runs[6].text = " vol. 1-2, pp. 19-25, 2015."
last_journal_para = p

p2 = paras[find_para(paras, " Kumar, Sarvesh")]
p2.runs[1].text = ("J. E. Stone et al., \u201cAccelerating molecular modeling applications with "
                    "graphics processors,\u201d")
p2.runs[2].text = " J. Comput. Chem.,"
p2.runs[3].text = " vol. 28, no. 16, pp. 2618-2640, 2007."

# All remaining sources are journal articles, not conference papers -- the
# template's "Conference:" sub-heading would misdescribe them, so it's
# cleared rather than left mislabeled above a list of journal citations.
set_text(paras[find_para(paras, "Conference:")], "")

p3 = paras[find_para(paras, "Salih, Azar")]
p3.runs[0].text = "J. A. Anderson, C. D. Lorenz, and A. Travesset, \u201cGeneral purpose molecular dynamics simulations fully implemented on graphics processing units,\u201d J. Comput. Phys., vol. 227, no. 10, pp. 5342-5359, 2008."
for r in p3.runs[1:]:
    r.text = ""

more_refs = [
    "P. Eastman et al., \u201cOpenMM 7: Rapid development of high performance algorithms for "
    "molecular dynamics,\u201d PLOS Comput. Biol., vol. 13, no. 7, 2017.",
    "J. Behler and M. Parrinello, \u201cGeneralized neural-network representation of "
    "high-dimensional potential-energy surfaces,\u201d Phys. Rev. Lett., vol. 98, 146401, 2007.",
    "K. T. Butler, D. W. Davies, H. Cartwright, O. Isayev, and A. Walsh, \u201cMachine learning "
    "for molecular and materials science,\u201d Nature, vol. 559, pp. 547-555, 2018.",
    "F. No\u00e9, A. Tkatchenko, K.-R. M\u00fcller, and C. Clementi, \u201cMachine learning for "
    "molecular simulation,\u201d Annu. Rev. Phys. Chem., vol. 71, pp. 361-390, 2020.",
    "J. White et al., \u201cA prompt pattern catalog to enhance prompt engineering with "
    "ChatGPT,\u201d arXiv:2302.11382, 2023.",
    "A. M. Bran, S. Cox, O. Schilter, C. Baldassari, A. D. White, and P. Schwaller, "
    "\u201cAugmenting large language models with chemistry tools,\u201d Nature Machine "
    "Intelligence, 2024.",
    "D. A. Boiko, R. MacKnight, B. Kline, and G. Gomes, \u201cAutonomous chemical research "
    "with large language models,\u201d Nature, vol. 624, pp. 570-578, 2023.",
]
anchor = p3
for ref in more_refs:
    anchor = clone_after(anchor, ref)

set_text(paras[find_para(paras, "A.P. Malvino and D.P. Leach")],
    "D. Frenkel and B. Smit, Understanding Molecular Simulation: From Algorithms to "
    "Applications, 2nd ed. San Diego, CA: Academic Press, 2002.")
paras = collect_paragraphs(doc)
book_para = paras[find_para(paras, "D. Frenkel and B. Smit")]
clone_after(book_para,
    "NVIDIA Corporation, CUDA C++ Programming Guide, Release 12.5, NVIDIA, 2024.")

doc.save(OUT)
print("Saved:", OUT)
