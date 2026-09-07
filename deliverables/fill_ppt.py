"""
Fills the Review-2 PPT template with real project content. Every slide's
placeholder text is edited in place (run-by-run, never text_frame.text=)
so the template's fonts/sizes/colors are preserved; new bullet lines are
added by cloning an existing paragraph's XML, matching the docx approach.

Slide 2 (Approval Mail From Guide) is left completely untouched -- it's a
process step (get the guide's email approval) that only the student can
complete, not project content.
"""
import copy
import os
from pptx import Presentation
from pptx.util import Inches, Pt

SRC = r"C:\Users\SHIVANI\Downloads\4 Review 2 PPT Template.pptx"
OUT = os.path.join(os.path.dirname(__file__), "Review2_PPT_Filled.pptx")
FIG = os.path.join(os.path.dirname(__file__), "figures")

prs = Presentation(SRC)
slides = prs.slides.__iter__()
slide_list = list(prs.slides)


def shape_by_ph_idx(slide, name_contains="Content Placeholder"):
    for shp in slide.shapes:
        if shp.has_text_frame and name_contains in shp.name:
            return shp
    return None


def shape_by_name(slide, name):
    for shp in slide.shapes:
        if shp.name == name:
            return shp
    return None


def set_run_text(paragraph, text):
    runs = paragraph.runs
    if not runs:
        paragraph.text = text
        return
    runs[0].text = text
    for r in runs[1:]:
        r.text = ""


def clone_paragraph_after(paragraph, text):
    """Clones a paragraph's XML (same run formatting) within the same
    text frame and inserts it right after, then sets its text."""
    new_p = copy.deepcopy(paragraph._p)
    paragraph._p.addnext(new_p)
    from pptx.text.text import _Paragraph
    new_para = _Paragraph(new_p, paragraph._parent)
    set_run_text(new_para, text)
    return new_para


def fill_bullets(shape, lines):
    """Fills a text frame's paragraphs with `lines`, reusing the
    existing (already-styled) paragraphs first and cloning the last
    template paragraph for any extra lines needed."""
    tf = shape.text_frame
    template_paras = list(tf.paragraphs)
    n_existing = len(template_paras)
    for i, line in enumerate(lines):
        if i < n_existing:
            set_run_text(template_paras[i], line)
        else:
            template_paras.append(clone_paragraph_after(template_paras[-1], line))
    # if the template had MORE paragraph slots than we needed, blank the rest
    for j in range(len(lines), n_existing):
        set_run_text(template_paras[j], "")


def slide(n):
    return slide_list[n - 1]


# ============================================================
# SLIDE 1 -- Title
# ============================================================
s = slide(1)
title_shape = None
team_shape = None
for shp in s.shapes:
    if not shp.has_text_frame:
        continue
    if "<<Project Title>>" in shp.text_frame.text:
        title_shape = shp
    elif "Team" in shp.text_frame.text and "Name" in shp.text_frame.text:
        team_shape = shp

for p in title_shape.text_frame.paragraphs:
    for r in p.runs:
        if r.text == "<<Programme ":
            r.text = "B.Tech (CSE)"
        elif r.text in ("B.Tech", "M.Tech", "/ "):
            r.text = ""
        elif "Course Code" in r.text:
            r.text = "  |  BCSE497J - Project-I"
        elif r.text == "<<Project Title>>":
            r.text = ("AI-Assisted Generation of CUDA-Based Molecular Dynamics "
                       "Simulations from Natural-Language Descriptions")

# The "Team / members / <<Name>> (<<Reg.No>>) / Faculty guide : <<Guidename>>"
# block lives in a separate (non-placeholder) text box on this slide.
name_done = False
for p in team_shape.text_frame.paragraphs:
    for r in p.runs:
        if r.text.startswith("<<Name >>") and not name_done:
            r.text = "Shivani Srivastava\t\t  ("
        elif r.text == "Reg.No" and not name_done:
            r.text = "23BDS0085"
        elif r.text == ">>)" and not name_done:
            r.text = ")"
            name_done = True  # only touch the FIRST (real) member slot
        elif r.text == "<<" :
            r.text = ""
        elif r.text == "Guidename":
            r.text = "Dr. Vijayasherly V"
        elif r.text == ">>":
            r.text = ""

# ============================================================
# SLIDE 3 -- Aim
# ============================================================
s = slide(3)
shp = shape_by_ph_idx(s)
set_run_text(shp.text_frame.paragraphs[0],
    "To build a local system where a user can describe, or manually configure, a molecular "
    "dynamics simulation and see it actually computed and visualized in real time \u2014 using "
    "a local, rule-based language component for parameter extraction and a hand-verified C++ "
    "simulation engine, with the design structured for future CUDA/GPU acceleration.")

# ============================================================
# SLIDE 4 -- Abstract
# ============================================================
s = slide(4)
shp = shape_by_ph_idx(s)
fill_bullets(shp, [
    "A local, browser-based tool that turns a plain-English description \u2014 or a manual "
    "form \u2014 into a real, physically simulated molecular dynamics run. A local rule-based "
    "language component extracts structured parameters (particle count, state of matter, "
    "interaction model, container, boundaries) without any external AI API.",
    "Validated parameters drive a compiled C++ engine implementing Lennard-Jones and "
    "Hard-Sphere models, real container/periodic boundaries, and genuine solid/liquid/gas "
    "behaviour. Real trajectories and energy data are shown in an interactive 3D viewer \u2014 "
    "CPU-based today, with the engine structured for a future CUDA/GPU port.",
])

# ============================================================
# SLIDE 5 -- Literature Review (empty placeholder -- add real bullets)
# ============================================================
s = slide(5)
content_shape = shape_by_name(s, "Content Placeholder 4")
# This placeholder is a generic "Content" box (picture/table/text) that was
# never populated in the template, so it has no <p:txBody> of its own yet --
# python-pptx's picture-flavored wrapper won't add one. Borrow the already-
# formatted txBody from another slide's matching content placeholder so this
# slide's bullet text uses the SAME font/size as every other content slide,
# instead of guessing formatting from scratch.
import copy as _copy
from pptx.text.text import _Paragraph
P_NS = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
A_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
donor_tx = shape_by_ph_idx(slide(6))._element.find(P_NS + "txBody")
new_tx = _copy.deepcopy(donor_tx)
content_shape._element.append(new_tx)
first_p_elem = new_tx.find(A_NS + "p")
p0 = _Paragraph(first_p_elem, content_shape)
set_run_text(p0, "Frenkel & Smit, Understanding Molecular Simulation \u2014 classical "
                  "Lennard-Jones potential, Verlet integration, periodic boundaries; "
                  "physical basis for this project's C++ engine.")
for extra in [
    "Plimpton (1995), LAMMPS \u2014 parallel/GPU molecular dynamics at scale, but requires "
    "manual configuration, no natural-language interface.",
    "Abraham et al. (2015), GROMACS \u2014 highly optimized MD, still needs manual scripting.",
    "NVIDIA CUDA C++ Programming Guide \u2014 reference for the engine's structure and its "
    "future GPU port.",
    "White et al. (2023) / OpenAI function calling \u2014 structured-output prompting, basis "
    "for restricting language extraction to a fixed parameter schema.",
]:
    p0 = clone_paragraph_after(p0, extra)
# the donor txBody had 2 template paragraphs; anything left over after our
# 5 real bullets (i.e. its original, still-unfilled 2nd paragraph) is blanked
for p_elem in list(new_tx.findall(A_NS + "p")):
    para = _Paragraph(p_elem, content_shape)
    if "<<" in para.text or "…" in para.text:
        set_run_text(para, "")

# ============================================================
# SLIDE 6 -- Research Gap
# ============================================================
s = slide(6)
shp = shape_by_ph_idx(s)
fill_bullets(shp, [
    "Existing MD packages (LAMMPS, GROMACS) are powerful but require manual configuration "
    "or direct code \u2014 no natural-language entry point exists.",
    "General-purpose AI coding assistants can generate code, but nothing guarantees the "
    "physics is correct. This project closes that gap: a language component restricted to "
    "parameter extraction only, paired with a fixed, hand-verified simulation engine.",
])

# ============================================================
# SLIDE 7 -- Objectives (+ SDG/Outcomes line on the second shape)
# ============================================================
s = slide(7)
shp = shape_by_ph_idx(s)
fill_bullets(shp, [
    "Design a structured parameter schema (particle count, density/temperature or speed, "
    "state, model, container/boundary).",
    "Build a local, rule-based natural-language extractor \u2014 no external AI service.",
    "Validate extracted or manual parameters against physically sensible ranges.",
    "Implement a C++ engine: Lennard-Jones + Hard-Sphere models, periodic/cuboid/spherical "
    "boundaries, solid/liquid/gas as real physics.",
    "Produce real trajectories and energy output; build an interactive 3D visualization.",
    "Structure the engine for a future CUDA/GPU port.",
])
for shp2 in s.shapes:
    if shp2.has_text_frame and "Identified" in shp2.text_frame.text:
        for p in shp2.text_frame.paragraphs:
            for r in p.runs:
                if r.text.startswith("Identified:"):
                    r.text = "Identified:  Sustainable Development Goals \u2014 SDG 9 (Primary), SDG 4 (Secondary)   |   "
                elif r.text == "Outcomes (":
                    r.text = "Outcomes: "
                elif "(Patent / Conference" in r.text:
                    r.text = "Product (Working Prototype)"
                elif "Journal (Scopus)" in r.text:
                    r.text = ""

# ============================================================
# SLIDE 8 -- Framework/Architecture/Block Diagram
# ============================================================
s = slide(8)
shp = shape_by_ph_idx(s)
fill_bullets(shp, [
    "The frontend collects a manual or natural-language request; the backend extracts and "
    "validates parameters (never touching simulation logic) and hands them to a compiled "
    "C++ engine, which writes real trajectory and energy files back for 3D visualization.",
])
s.shapes.add_picture(os.path.join(FIG, "fig2_architecture.png"),
                      left=Inches(1.3), top=Inches(2.6), width=Inches(7.4))

# ============================================================
# SLIDE 9 -- Functional Requirements
# ============================================================
s = slide(9)
shp = shape_by_ph_idx(s)
fill_bullets(shp, [
    "Accept a natural-language description and extract structured parameters locally.",
    "Accept manually entered parameters as an equivalent alternative input path.",
    "Validate all parameters against physical ranges before simulation, with clear errors.",
    "Run validated parameters through the compiled C++ engine (Lennard-Jones / Hard-Sphere).",
    "Support periodic boundaries and real cuboidal/spherical containers.",
    "Represent solid, liquid, and gas as genuine differences in arrangement and motion.",
    "Record real particle trajectories and energy, and display them in an interactive 3D "
    "view with an energy-vs-time chart.",
])

# ============================================================
# SLIDE 10 -- Modules
# ============================================================
s = slide(10)
shp = shape_by_ph_idx(s)
fill_bullets(shp, [
    "Frontend (HTML/CSS/JS): Setup \u2192 Review \u2192 Run workflow, 3D viewer, energy chart.",
    "Local NLP Extractor: rule-based parsing of natural-language simulation descriptions.",
    "Parameter Validator: shared range/feasibility checks (manual and NL input alike).",
    "C++ Simulation Engine: Lennard-Jones & Hard-Sphere models, boundaries, states of matter.",
    "Trajectory & Energy Output: real particle positions and energy written per run.",
    "3D Visualization: hand-built renderer with rotate/zoom, playing back real trajectories.",
])

# ============================================================
# SLIDE 11 -- Experiments and Results
# ============================================================
s = slide(11)
shp = shape_by_ph_idx(s)
fill_bullets(shp, [
    "21 automated tests (pytest) covering validation rules, NLP extraction, and engine "
    "physics all pass.",
    "Periodic (NVE) runs conserve total energy to within 0.5 reduced units over a full run.",
    "Hard-Sphere elastic collisions conserve kinetic energy almost exactly (drift < 0.05).",
    "The liquid state visibly settles into the bottom ~30% of a walled container, confirmed "
    "in both automated tests and a live end-to-end browser run.",
    "The full pipeline \u2014 natural-language input through 3D visualization \u2014 was verified "
    "live: real parameters, a real C++ run, and a real rendered trajectory.",
])

# ============================================================
# SLIDE 12 -- Conclusion
# ============================================================
s = slide(12)
shp = shape_by_ph_idx(s)
fill_bullets(shp, [
    "This project delivers a working, locally-run system that turns natural-language or "
    "manual input into a real, physically simulated and visualized molecular dynamics run.",
    "Restricting the language component to parameter extraction \u2014 never simulation logic "
    "\u2014 keeps the physics fixed, hand-verified, and trustworthy.",
    "The current CPU implementation is deliberately structured so its core routines can be "
    "ported to CUDA/GPU acceleration as future work, without redesigning the system.",
])

# ============================================================
# SLIDE 13 -- References
# ============================================================
s = slide(13)
for shp in s.shapes:
    if not shp.has_text_frame:
        continue
    for p in shp.text_frame.paragraphs:
        for r in p.runs:
            if r.text.startswith('[1] J. K. Author'):
                r.text = 'S. Plimpton, "Fast parallel algorithms for short-range molecular dynamics," '
            elif r.text == "Name of Journal":
                r.text = "Journal of Computational Physics"
            elif r.text.startswith(", vol. 15, no. 3"):
                r.text = ", vol. 117, no. 1, pp. 1-19, 1995."
            elif r.text.startswith('[2] A. B. Researcher'):
                r.text = 'NVIDIA Corporation, '
            elif r.text == "Proc. Name of Conference":
                r.text = "CUDA C++ Programming Guide"
            elif r.text.startswith(", City, Country"):
                r.text = ", Release 12.5, NVIDIA, 2024."

doc_last = [p for shp in s.shapes if shp.has_text_frame for p in shp.text_frame.paragraphs
            if p.runs and "Release 12.5" in "".join(r.text for r in p.runs)]
if doc_last:
    clone_paragraph_after(doc_last[0],
        'M. J. Abraham et al., "GROMACS: High performance molecular simulations through '
        'multi-level parallelism from laptops to supercomputers," SoftwareX, vol. 1-2, '
        'pp. 19-25, 2015.')

prs.save(OUT)
print("Saved:", OUT)
